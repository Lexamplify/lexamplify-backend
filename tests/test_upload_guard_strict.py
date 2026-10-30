"""Unit tests for the stricter upload helper in utils/upload_guard.py (Step 7, batch 2a).

The helper sits on top of check_upload and is used by /api/ai/extract-file, /api/upload-case, /api/vc/simulate and the two
conflict routes. check_upload itself, and so the contract routes, must behave exactly as before; test_check_upload_is_unchanged
pins that. Every file here is built in memory or is an existing fixture; nothing touches the network or real data."""
import io
import time
import zipfile
from pathlib import Path

import pytest

from utils import upload_guard as ug

FX = Path(__file__).resolve().parent / "fixtures" / "indian_legal"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
      '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
      '<Default Extension="xml" ContentType="application/xml"/>'
      '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')


# ── builders ────────────────────────────────────────────────────────────────────────────────────
def doc_xml(unit, n, prefix="w", head_extra="", before_root=""):
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>{before_root}<{prefix}:document xmlns:{prefix}="{W}"{head_extra}>'
            f'<{prefix}:body>{unit * n}</{prefix}:body></{prefix}:document>')


def para(prefix="w", text="Lorem ipsum dolor sit amet."):
    return f'<{prefix}:p><{prefix}:r><{prefix}:t>{text}</{prefix}:t></{prefix}:r></{prefix}:p>'


def varied_xml(n, prefix="w", shape="para"):
    """A main part of n paragraphs (or n one-cell tables) whose text differs every time, so the compression-ratio rule does
    not fire and the element-count rule is what is being tested."""
    import hashlib
    unit = (lambda t: para(prefix, t)) if shape == "para" else \
        (lambda t: f'<{prefix}:tbl><{prefix}:tr><{prefix}:tc><{prefix}:p><{prefix}:r><{prefix}:t>{t}</{prefix}:t></{prefix}:r></{prefix}:p></{prefix}:tc></{prefix}:tr></{prefix}:tbl>')
    body = "".join(unit(hashlib.md5(str(i).encode()).hexdigest() * 2) for i in range(n))
    return f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><{prefix}:document xmlns:{prefix}="{W}"><{prefix}:body>{body}</{prefix}:body></{prefix}:document>'


def zip_docx(document_xml, extra=None, rels=RELS, level=9):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=level) as z:
        z.writestr("[Content_Types].xml", CT)
        if rels is not None:
            z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document_xml)
        for name, blob in (extra or {}).items():
            z.writestr(name, blob)
    return buf.getvalue()


def pdf_with_text(text="The appellant submits that the learned court erred in law.", pages=1):
    import fitz
    d = fitz.open()
    for _ in range(pages):
        d.new_page().insert_text((72, 72), text)
    b = d.tobytes()
    d.close()
    return b


def pdf_scanned():
    """A page that is only a picture (no text layer), like a scan."""
    import fitz
    d = fitz.open()
    p = d.new_page()
    p.insert_image(fitz.Rect(50, 50, 300, 300), pixmap=fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 60), False))
    b = d.tobytes()
    d.close()
    return b


def pdf_password():
    import fitz
    d = fitz.open()
    d.new_page().insert_text((72, 72), "secret")
    b = d.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    d.close()
    return b


def pdf_owner_only():
    """Encrypted with an owner password only: opens with no password (many court PDFs are like this)."""
    import fitz
    d = fitz.open()
    d.new_page().insert_text((72, 72), "owner password only")
    b = d.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="")
    d.close()
    return b


def code_of(problem):
    return problem["code"] if problem else None


# ── check_upload must not change ────────────────────────────────────────────────────────────────
def test_check_upload_is_unchanged():
    import inspect
    assert str(inspect.signature(ug.check_upload)) == "(data: bytes, filename: str)"
    assert (ug.MAX_UPLOAD_BYTES, ug.MAX_UNCOMPRESSED, ug.MAX_PART, ug.MAX_RATIO, ug.MAX_ENTRIES, ug.MAX_PDF_PAGES) == \
        (30 * 1024 * 1024, 200 * 1024 * 1024, 120 * 1024 * 1024, 400, 4000, 1500)
    assert ug.check_upload(b"hello", "a.txt") == ("txt", None)
    assert ug.check_upload(b"hello", "a.md")[1]["code"] == "UNSUPPORTED_FORMAT"        # the strict helper must not rely on this changing
    assert ug.check_upload(b"", "a.pdf")[1]["code"] == "EMPTY_FILE"
    assert ug.check_upload(b"x", "a.pdf")[1] == {"code": "NOT_A_PDF", "message": ug._MESSAGES["NOT_A_PDF"], "status": 422}
    assert ug.check_upload(b"x" * (30 * 1024 * 1024 + 1), "a.txt")[1]["status"] == 413


# ── the element counter ─────────────────────────────────────────────────────────────────────────
def _independent_count(xml):
    from lxml import etree
    root = etree.fromstring(xml)
    return sum(1 for e in root.iter() if isinstance(e.tag, str) and e.tag in (f"{{{W}}}p", f"{{{W}}}r", f"{{{W}}}tc"))


DOCX_FIXTURES = sorted(p for p in FX.glob("*.docx") if p.name not in ("macro_enabled.docx", "old_word_renamed.docx"))


@pytest.mark.parametrize("path", DOCX_FIXTURES, ids=lambda p: p.name)
def test_counter_matches_an_independent_count_on_real_fixtures(path):
    zf = zipfile.ZipFile(path)
    xml = zf.read("word/document.xml")
    n, status = ug.count_docx_elements(xml, 25000)
    assert status == "ok" and n == _independent_count(xml) and n > 0


def test_counter_ignores_the_prefix_and_counts_by_namespace():
    xml = doc_xml(para("x"), 15000, prefix="x").encode()          # 15,000 x 2 = 30,000 elements, bound to the Word namespace as "x"
    t = time.time()
    n, status = ug.count_docx_elements(xml, 25000)
    assert status == "limit" and n == 25001 and time.time() - t < 2


def test_counter_does_not_count_other_namespaces_or_other_elements():
    other = '<a:p xmlns:a="urn:other"><a:r/></a:p>' * 50000        # same local names, a different namespace
    xml = (f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>{other}<w:p><w:pPr/><w:r><w:rPr/><w:t>x</w:t></w:r></w:p></w:body></w:document>').encode()
    assert ug.count_docx_elements(xml, 25000) == (2, "ok")


def test_counter_stops_as_soon_as_the_cap_is_passed():
    xml = doc_xml(para(), 300000).encode()                        # 600,000 counted elements, about 18 MB of xml
    t = time.time()
    n, status = ug.count_docx_elements(xml, 1000)
    assert (n, status) == (1001, "limit") and time.time() - t < 1


def test_counter_rejects_a_doctype_and_entity_declarations():
    xml = doc_xml(para(), 3, before_root='<!DOCTYPE w:document [<!ENTITY a "bbbbbbbbbb">]>').encode()
    assert ug.count_docx_elements(xml, 25000)[1] == "dtd"
    xml2 = doc_xml(para(), 3, before_root='<!DOCTYPE w:document>').encode()
    assert ug.count_docx_elements(xml2, 25000)[1] == "dtd"
    # also for a document with no counted element at all
    xml3 = (f'<?xml version="1.0"?><!DOCTYPE w:document [<!ENTITY a "b">]><w:document xmlns:w="{W}"/>').encode()
    assert ug.count_docx_elements(xml3, 25000)[1] == "dtd"


def test_counter_does_not_expand_entities():
    lol = '<!DOCTYPE lolz [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;"><!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">]>'
    xml = (f'<?xml version="1.0"?>{lol}<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>&c;&c;&c;&c;</w:t></w:r></w:p></w:body></w:document>').encode()
    t = time.time()
    assert ug.count_docx_elements(xml, 25000)[1] == "dtd" and time.time() - t < 1


@pytest.mark.parametrize("xml", [b"", b"not xml at all", b"<w:document", b'<?xml version="1.0"?><w:document xmlns:w="%s"><w:body><w:p></w:body></w:document>' % W.encode()])
def test_counter_reports_malformed_xml(xml):
    assert ug.count_docx_elements(xml, 25000)[1] == "malformed"


# ── docx rules ──────────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("path", DOCX_FIXTURES, ids=lambda p: p.name)
def test_every_valid_docx_fixture_passes(path):
    assert ug.check_upload_strict(path.read_bytes(), path.name) == ("docx", None)


def test_macro_and_renamed_old_word_fixtures_are_still_refused():
    assert code_of(ug.check_upload_strict((FX / "macro_enabled.docx").read_bytes(), "a.docx")[1]) == "MACRO_DOCUMENT"
    assert code_of(ug.check_upload_strict((FX / "old_word_renamed.docx").read_bytes(), "a.docx")[1]) == "NOT_A_DOCX"


def test_compressible_docx_bomb_is_refused_by_the_ratio_rule():
    data = zip_docx(doc_xml(para(), 210000))                       # about 12 MB of xml in a tiny zip
    assert len(data) < 100 * 1024
    t = time.time()
    kind, problem = ug.check_upload_strict(data, "bomb.docx")
    assert kind is None and problem["code"] == "UNSAFE_ARCHIVE" and problem["status"] == 422 and time.time() - t < 3
    assert ug.check_upload(data, "bomb.docx") == ("docx", None), "the old guard lets this one through (that is the gap)"


def test_ratio_rule_applies_to_any_part_above_256kb():
    data = zip_docx(doc_xml(para(), 3), extra={"word/media/blob.bin": b"\x00" * (3 * 1024 * 1024)})
    assert code_of(ug.check_upload_strict(data, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_ratio_rule_does_not_touch_small_parts():
    data = zip_docx(doc_xml(para(), 3), extra={"word/media/small.bin": b"\x00" * (200 * 1024)})
    assert ug.check_upload_strict(data, "a.docx") == ("docx", None)


def test_per_part_cap(monkeypatch):
    monkeypatch.setattr(ug, "STRICT_DOCX_MAX_PART", 50 * 1024)
    import os
    data = zip_docx(doc_xml(para(), 3), extra={"word/media/photo.bin": os.urandom(80 * 1024)})
    assert code_of(ug.check_upload_strict(data, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_total_uncompressed_cap(monkeypatch):
    import os
    monkeypatch.setattr(ug, "STRICT_DOCX_MAX_TOTAL", 100 * 1024)
    data = zip_docx(doc_xml(para(), 3), extra={"word/media/a.bin": os.urandom(60 * 1024), "word/media/b.bin": os.urandom(60 * 1024)})
    assert code_of(ug.check_upload_strict(data, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_default_limits_have_the_agreed_values():
    assert ug.STRICT_DOCX_MAX_PART == 8 * 1024 * 1024
    assert ug.STRICT_DOCX_MAX_TOTAL == 24 * 1024 * 1024
    assert ug.STRICT_DOCX_MAX_RATIO == 200 and ug.STRICT_DOCX_RATIO_FROM == 256 * 1024
    assert ug.STRICT_DOCX_MAX_ELEMENTS == 25000
    assert ug.STRICT_MAX_PDF_PAGES == 1500
    assert ug.CONFLICT_MAX_FILES == 10 and ug.CONFLICT_MAX_TOTAL_BYTES == 50 * 1024 * 1024
    assert ug.CONFLICT_MAX_DOC_CHARS == 600000


def test_docx_with_an_other_prefix_and_too_many_elements_is_refused_quickly():
    data = zip_docx(varied_xml(15000, "x"), level=1)
    t = time.time()
    kind, problem = ug.check_upload_strict(data, "alt.docx")
    assert kind is None and problem["code"] == "UNSAFE_ARCHIVE" and time.time() - t < 3
    assert "25,000" in problem["message"]


def test_docx_with_too_many_table_cells_is_refused():
    data = zip_docx(varied_xml(9000, shape="cell"))                           # 3 elements per cell: 27,000
    assert code_of(ug.check_upload_strict(data, "t.docx")[1]) == "UNSAFE_ARCHIVE"


def test_docx_at_just_under_the_cap_is_accepted():
    data = zip_docx(varied_xml(12000))                        # 24,000 elements
    assert ug.check_upload_strict(data, "ok.docx") == ("docx", None)


def test_docx_with_doctype_or_malformed_main_part():
    bad = zip_docx(doc_xml(para(), 3, before_root='<!DOCTYPE w:document [<!ENTITY a "b">]>'))
    assert code_of(ug.check_upload_strict(bad, "a.docx")[1]) == "UNSAFE_ARCHIVE"
    broken = zip_docx(doc_xml(para(), 3)[:-30])
    assert code_of(ug.check_upload_strict(broken, "a.docx")[1]) == "NOT_A_DOCX"


def test_docx_without_valid_package_relationships_is_not_a_docx():
    assert code_of(ug.check_upload_strict(zip_docx(doc_xml(para(), 3), rels=None), "a.docx")[1]) == "NOT_A_DOCX"
    wrong = RELS.replace("word/document.xml", "word/missing.xml")
    assert code_of(ug.check_upload_strict(zip_docx(doc_xml(para(), 3), rels=wrong), "a.docx")[1]) == "NOT_A_DOCX"


def test_the_counted_part_is_the_one_named_in_the_package_relationships():
    """A docx whose relationships point at another part is judged on that part, not on a decoy word/document.xml."""
    rels = RELS.replace("word/document.xml", "word/main.xml")
    data = zip_docx(doc_xml(para(), 3), extra={"word/main.xml": varied_xml(20000)}, rels=rels)
    assert code_of(ug.check_upload_strict(data, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_headers_and_footers_are_not_read_by_the_extractors():
    """Why only the main part is counted: python-docx text extraction never iterates a footer, so a big one costs only its parse."""
    import hashlib
    footer = f'<?xml version="1.0"?><w:ftr xmlns:w="{W}">' + "".join(para(text=hashlib.md5(str(i).encode()).hexdigest() * 2) for i in range(60000)) + "</w:ftr>"
    data = zip_docx(doc_xml(para(), 3), extra={"word/footer1.xml": footer})
    # accepted by the helper, and the real extraction stays fast:
    assert ug.check_upload_strict(data, "a.docx") == ("docx", None)


def test_realistic_300_page_contract_passes_all_limits():
    """About 300 pages: 8-9 paragraphs a page with 3-4 runs each, and one page in five carries a table (8 rows x 4 columns)."""
    import docx
    d = docx.Document()
    for page in range(300):
        if page % 5 == 4:
            t = d.add_table(rows=8, cols=4)
            for row in t.rows:
                for cell in row.cells:
                    cell.paragraphs[0].add_run("Schedule item ")
                    cell.paragraphs[0].add_run("Rs. 4,75,000").bold = True
            for _ in range(3):
                p = d.add_paragraph()
                for r in range(3):
                    p.add_run(f"Payment shall be made within thirty days of invoice {r}. ")
        else:
            for k in range(8 + page % 2):
                p = d.add_paragraph()
                p.add_run(f"{k + 1}. ")
                p.add_run("The Company", ).bold = True
                p.add_run(" shall indemnify the Contractor against all losses arising from breach of this Agreement. ")
                if k % 2:
                    p.add_run("Disputes are subject to the courts at Mumbai.")
        d.add_page_break()
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    n, status = ug.count_docx_elements(xml, 10 ** 9)
    print(f"\nREALISTIC 300-page docx: file={len(data)} bytes, document.xml={len(xml)} bytes, counted elements={n} (cap {ug.STRICT_DOCX_MAX_ELEMENTS})")
    assert n < ug.STRICT_DOCX_MAX_ELEMENTS
    assert ug.check_upload_strict(data, "contract.docx") == ("docx", None)


# ── pdf rules ───────────────────────────────────────────────────────────────────────────────────
PDF_FIXTURES = sorted(p for p in FX.glob("*.pdf") if p.name not in ("empty.pdf", "fake_pdf.pdf", "password_protected.pdf", "truncated.pdf"))


@pytest.mark.parametrize("path", PDF_FIXTURES, ids=lambda p: p.name)
def test_every_valid_pdf_fixture_passes(path):
    assert ug.check_upload_strict(path.read_bytes(), path.name) == ("pdf", None)


def test_bad_pdf_fixtures_are_refused_with_the_contract_codes():
    for name, code in (("empty.pdf", "EMPTY_FILE"), ("fake_pdf.pdf", "NOT_A_PDF"), ("truncated.pdf", "NOT_A_PDF"), ("password_protected.pdf", "EXTRACTION_ERROR")):
        kind, problem = ug.check_upload_strict((FX / name).read_bytes(), name)
        assert kind is None and problem["code"] == code and problem["status"] == 422, name


def test_scanned_pdf_without_any_text_is_accepted():
    data = pdf_scanned()
    import fitz
    d = fitz.open(stream=data, filetype="pdf")
    assert d.page_count == 1 and not d[0].get_text().strip()
    assert ug.check_upload_strict(data, "scan.pdf") == ("pdf", None)
    assert ug.check_upload_strict((FX / "scanned_contract.pdf").read_bytes(), "scan.pdf") == ("pdf", None)


def test_password_pdf_is_refused_but_owner_password_only_pdf_is_accepted():
    kind, problem = ug.check_upload_strict(pdf_password(), "p.pdf")
    assert kind is None and problem["code"] == "EXTRACTION_ERROR" and problem["status"] == 422 and "password" in problem["message"].lower()
    assert ug.check_upload_strict(pdf_owner_only(), "o.pdf") == ("pdf", None)


def test_junk_with_a_pdf_header_is_refused_fast():
    junk = b"%PDF-1.4\n" + b"\x00" * (1024 * 1024 - 9)
    t = time.time()
    kind, problem = ug.check_upload_strict(junk, "x.pdf")
    assert kind is None and problem["code"] == "NOT_A_PDF" and time.time() - t < 1
    objs = b"%PDF-1.4\n" + (b"1 0 obj\n<< /Type /Catalog >>\nendobj\n" * 30000)
    assert code_of(ug.check_upload_strict(objs, "x.pdf")[1]) == "NOT_A_PDF"


def test_pdf_page_limit(monkeypatch):
    monkeypatch.setattr(ug, "STRICT_MAX_PDF_PAGES", 2)
    assert ug.check_upload_strict(pdf_with_text(pages=2), "a.pdf") == ("pdf", None)
    kind, problem = ug.check_upload_strict(pdf_with_text(pages=3), "a.pdf")
    assert kind is None and problem["code"] == "TOO_MANY_PAGES" and problem["status"] == 413
    assert "2" in problem["message"]


def test_broken_page_trees_are_refused_without_a_crash():
    def pdf(objs):
        out, offs = b"%PDF-1.4\n", []
        for i, body in enumerate(objs, 1):
            offs.append(len(out))
            out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
        x = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
        return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x)
    loop = pdf([b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [2 0 R] /Count 1 >>"])
    huge = pdf([b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 999999999 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 9 9] >>"])
    for data in (loop, huge):
        t = time.time()
        kind, problem = ug.check_upload_strict(data, "x.pdf")
        assert kind is None and problem["code"] == "NOT_A_PDF" and time.time() - t < 2


# ── names, bytes and text ───────────────────────────────────────────────────────────────────────
def test_empty_files_are_refused_whatever_the_name():
    for name in ("a.pdf", "a.docx", "a.txt", "a.md", "noextension"):
        assert code_of(ug.check_upload_strict(b"", name)[1]) == "EMPTY_FILE"


def test_text_files_of_any_extension_are_accepted_as_text():
    hindi = "यह एक सरल पाठ फ़ाइल है। नोटिस ₹4,75,000".encode("utf-8")
    for name in ("notes.txt", "README.md", "table.csv", "page.html", "data.json", "noextension", "x.PDF.exe", "a.TXT"):
        assert ug.check_upload_strict(b"Plain text of a notice. Payment within 30 days.", name) == ("txt", None), name
    assert ug.check_upload_strict(hindi, "notice.md") == ("txt", None)
    assert ug.check_upload_strict("hello".encode("utf-16"), "u16.txt") == ("txt", None)         # with a byte-order mark


def test_binary_data_under_a_text_name_is_refused():
    assert code_of(ug.check_upload_strict(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR" + b"\x00" * 100, "pic.png")[1]) == "NOT_TEXT"
    assert code_of(ug.check_upload_strict(b"\x00\x01\x02" * 100, "a.txt")[1]) == "NOT_TEXT"


def test_strict_mode_refuses_text_named_pdf_or_docx():
    assert code_of(ug.check_upload_strict(b"this is plain text pretending to be a pdf", "a.pdf")[1]) == "NOT_A_PDF"
    assert code_of(ug.check_upload_strict(b"this is plain text pretending to be a docx", "a.docx")[1]) == "NOT_A_DOCX"


def test_conflict_mode_accepts_clean_text_under_a_document_name():
    kw = {"allow_text_under_document_name": True}
    assert ug.check_upload_strict(b"Payment within 30 days. Courts at Delhi.", "restored.pdf", **kw) == ("txt", None)
    assert ug.check_upload_strict("भुगतान 30 दिनों में".encode(), "restored.docx", **kw) == ("txt", None)
    # ... but only clean text: not NUL bytes, not something that starts like a PDF or a zip
    assert code_of(ug.check_upload_strict(b"abc\x00def", "a.pdf", **kw)[1]) == "NOT_A_PDF"
    assert code_of(ug.check_upload_strict(b"abc\x00def", "a.docx", **kw)[1]) == "NOT_A_DOCX"
    assert code_of(ug.check_upload_strict(b"%PDF-1.4 but nothing else", "a.pdf", **kw)[1]) == "NOT_A_PDF"
    assert code_of(ug.check_upload_strict(b"PK not really a zip", "a.docx", **kw)[1]) == "NOT_A_DOCX"
    assert code_of(ug.check_upload_strict(b"", "a.pdf", **kw)[1]) == "EMPTY_FILE"


def test_conflict_mode_still_judges_real_pdfs_and_docx_by_their_bytes():
    kw = {"allow_text_under_document_name": True}
    assert code_of(ug.check_upload_strict(pdf_password(), "a.pdf", **kw)[1]) == "EXTRACTION_ERROR"
    assert code_of(ug.check_upload_strict(zip_docx(doc_xml(para(), 210000)), "a.docx", **kw)[1]) == "UNSAFE_ARCHIVE"
    assert code_of(ug.check_upload_strict((FX / "macro_enabled.docx").read_bytes(), "a.docx", **kw)[1]) == "MACRO_DOCUMENT"
    assert ug.check_upload_strict(pdf_with_text(), "a.pdf", **kw) == ("pdf", None)


def test_size_limit_uses_the_helper_constant(monkeypatch):
    monkeypatch.setattr(ug, "STRICT_MAX_BYTES", 1000)
    kind, problem = ug.check_upload_strict(b"a" * 1001, "a.txt")
    assert kind is None and problem["code"] == "TOO_LARGE" and problem["status"] == 413
    assert ug.check_upload_strict(b"a" * 1000, "a.txt") == ("txt", None)


class _CountingStream(io.BytesIO):
    def __init__(self, data):
        super().__init__(data)
        self.read_calls = []

    def read(self, n=-1):
        self.read_calls.append(n)
        return super().read(n)


class _Upload:
    def __init__(self, data, filename="a.txt"):
        self.stream = _CountingStream(data)
        self.filename = filename


def test_read_checked_upload_reads_at_most_max_plus_one_and_rewinds(monkeypatch):
    monkeypatch.setattr(ug, "STRICT_MAX_BYTES", 5000)
    up = _Upload(b"a" * 100000)
    kind, data, problem = ug.read_checked_upload(up)
    assert kind is None and data is None and problem["code"] == "TOO_LARGE"
    assert up.stream.read_calls == [5001] and up.stream.tell() == 0, up.stream.read_calls
    ok = _Upload(b"hello there", "n.md")
    kind, data, problem = ug.read_checked_upload(ok)
    assert (kind, data, problem) == ("txt", b"hello there", None) and ok.stream.tell() == 0


def test_read_checked_upload_passes_the_conflict_flag():
    up = _Upload(b"plain text", "restored.pdf")
    assert ug.read_checked_upload(up)[2]["code"] == "NOT_A_PDF"
    assert ug.read_checked_upload(_Upload(b"plain text", "restored.pdf"), allow_text_under_document_name=True)[0] == "txt"


def test_file_names_are_never_used_or_echoed():
    for name in ("../../x.pdf", "a\x00b.pdf", "C:\\secret\\client_name.pdf", "n" * 300 + ".pdf"):
        kind, problem = ug.check_upload_strict(b"just text", name)
        assert problem and "secret" not in problem["message"] and "client_name" not in problem["message"] and ".." not in problem["message"]


def test_every_message_is_plain_english_and_has_no_internal_detail():
    import re
    leak = re.compile(r"Traceback|\.py|sqlite|Exception|[A-Z]:\\|/Users/|word/|xml", re.I)
    samples = [ug.check_upload_strict(b"x", "a.pdf")[1], ug.check_upload_strict(pdf_password(), "a.pdf")[1],
               ug.check_upload_strict(zip_docx(doc_xml(para(), 210000)), "a.docx")[1], ug.check_upload_strict(zip_docx(varied_xml(15000, "x")), "a.docx")[1]]
    for p in samples:
        assert not leak.search(p["message"]), p["message"]
        assert p["message"].endswith(".") and len(p["message"]) > 40


def test_problem_json_shapes():
    p = {"code": "NOT_A_PDF", "message": "msg.", "status": 422}
    assert ug.problem_json(p) == {"error": "msg.", "code": "NOT_A_PDF"}
    assert ug.problem_json(p, with_message=True) == {"error": "msg.", "message": "msg.", "code": "NOT_A_PDF"}


# ── conflict request limits ─────────────────────────────────────────────────────────────────────
def _files(n, size=10, names=None):
    from werkzeug.datastructures import FileStorage, MultiDict
    md = MultiDict()
    for i in range(n):
        key = names[i] if names else f"doc{i + 1}"
        md.add(key, FileStorage(stream=io.BytesIO(b"x" * size), filename=f"f{i}.txt"))
    return md


def test_file_count_cap_counts_every_part_whatever_its_field_name():
    assert ug.check_files_request(_files(10)) is None
    for names in (None, [f"weird{i}" for i in range(11)], ["files"] * 11, ["documents"] * 5 + ["doc1", "zzz", "file", "a", "b", "c"]):
        p = ug.check_files_request(_files(11, names=names))
        assert p and p["code"] == "TOO_MANY_FILES" and p["status"] == 413 and "10" in p["message"]
    # a gap in the doc1..docN slots does not hide files from the count
    p = ug.check_files_request(_files(11, names=["doc2", "doc3", "doc4", "doc5", "doc6", "doc7", "doc8", "doc9", "doc10", "doc11", "doc12"]))
    assert p and p["code"] == "TOO_MANY_FILES"


def test_total_size_cap(monkeypatch):
    monkeypatch.setattr(ug, "CONFLICT_MAX_TOTAL_BYTES", 1000)
    assert ug.check_files_request(_files(4, size=250)) is None
    p = ug.check_files_request(_files(4, size=251))
    assert p and p["code"] == "TOO_LARGE" and p["status"] == 413 and "together" in p["message"].lower()


def test_each_file_in_a_conflict_request_gets_the_strict_checks_and_is_rewound():
    from werkzeug.datastructures import FileStorage, MultiDict
    md = MultiDict()
    md.add("doc1", FileStorage(stream=io.BytesIO(pdf_with_text()), filename="a.pdf"))
    md.add("doc2", FileStorage(stream=io.BytesIO(b"restored text under a pdf name"), filename="b.pdf"))
    assert ug.check_files_request(md) is None
    assert md["doc1"].stream.tell() == 0 and md["doc2"].read() == b"restored text under a pdf name"
    md.add("doc3", FileStorage(stream=io.BytesIO(pdf_password()), filename="c.pdf"))
    assert ug.check_files_request(md)["code"] == "EXTRACTION_ERROR"
    md2 = MultiDict()
    md2.add("doc1", FileStorage(stream=io.BytesIO(b""), filename="empty.txt"))
    assert ug.check_files_request(md2)["code"] == "EMPTY_FILE"


def test_parts_without_a_file_name_are_not_judged_but_are_counted():
    from werkzeug.datastructures import FileStorage, MultiDict
    md = MultiDict()
    for i in range(11):
        md.add(f"x{i}", FileStorage(stream=io.BytesIO(b""), filename=""))
    assert ug.check_files_request(md)["code"] == "TOO_MANY_FILES"
    md = MultiDict()
    md.add("x", FileStorage(stream=io.BytesIO(b""), filename=""))
    assert ug.check_files_request(md) is None


def _text(chars, line=99):
    """`chars` characters in lines of `line` characters, so only the size rules (not the longest-line rule) are in play."""
    return (("a" * line + "\n") * (chars // (line + 1) + 1))[:chars]


def _docs(n, chars):
    return [{"name": f"d{i}", "text": _text(chars)} for i in range(n)]


def test_text_length_caps_never_cut_silently():
    assert ug.check_conflict_documents(_docs(2, 600000)) is None
    p = ug.check_conflict_documents(_docs(2, 600001))
    assert p["code"] == "TOO_LARGE" and p["status"] == 413 and "600,000" in p["message"]
    p = ug.check_conflict_documents(_docs(11, 10))
    assert p["code"] == "TOO_MANY_FILES"


def test_pair_budget_allows_five_documents_of_300000_characters():
    assert ug.check_conflict_documents(_docs(5, 300000)) is None       # the "normal big conflict check" (guard rule 1)
    assert ug.check_conflict_documents(_docs(5, 300001))["code"] == "TOO_LARGE"
    assert ug.check_conflict_documents(_docs(2, 600000)) is None
    assert ug.check_conflict_documents(_docs(3, 600000)) is None
    assert ug.check_conflict_documents(_docs(4, 600000))["code"] == "TOO_LARGE"
    assert ug.check_conflict_documents(_docs(10, 66666)) is None
    assert ug.check_conflict_documents(_docs(10, 66667))["code"] == "TOO_LARGE"
    p = ug.check_conflict_documents(_docs(10, 600000))
    assert p["code"] == "TOO_LARGE" and p["status"] == 413 and "too much text to compare" in p["message"]


def test_pair_budget_is_inclusive_and_counts_every_document(monkeypatch):
    monkeypatch.setattr(ug, "CONFLICT_MAX_PAIR_CHARS", 1000)
    two = lambda n: [{"name": "a", "text": "x" * n}, {"name": "b", "text": "y" * n}]          # (2 - 1) x 2n
    assert ug.check_conflict_documents(two(500)) is None                                       # exactly 1,000: accepted
    assert ug.check_conflict_documents(two(501))["code"] == "TOO_LARGE"                         # 1,002: refused
    three = [{"name": str(i), "text": "x" * 250} for i in range(3)]                             # (3 - 1) x 750 = 1,500
    assert ug.check_conflict_documents(three)["code"] == "TOO_LARGE"
    three = [{"name": str(i), "text": "x" * 166} for i in range(3)]                             # (3 - 1) x 498 = 996
    assert ug.check_conflict_documents(three) is None


def test_longest_line_rule_is_inclusive_and_splits_on_newline_only(monkeypatch):
    monkeypatch.setattr(ug, "CONFLICT_MAX_LINE_CHARS", 50)
    ok = {"name": "ok", "text": "x" * 50 + "\n" + "y" * 10}
    assert ug.check_conflict_documents([ok, ok]) is None                                       # exactly 50 characters between newlines
    over = {"name": "over", "text": "x" * 51 + "\n" + "y" * 10}
    p = ug.check_conflict_documents([ok, over])
    assert p["code"] == "TOO_LARGE" and p["status"] == 413 and "20,000" not in p["message"] and "50" in p["message"]
    assert "over" not in p["message"].replace("without", "")                                   # no document name in the message
    # other line separators do not count as line breaks (the comparator does not use them either)
    for sep in ("\r", " ", " ", "\x0b", "\x0c", "\x85", "\r\r"):
        text = (("x" * 20) + sep) * 5                                                           # 105+ characters, no "\n"
        assert ug.check_conflict_documents([ok, {"name": "s", "text": text}])["code"] == "TOO_LARGE", repr(sep)
    # a long document made of short lines is fine
    assert ug.check_conflict_documents([ok, {"name": "s", "text": ("x" * 40 + "\n") * 200}]) is None


def test_long_line_rule_default_and_message():
    assert ug.CONFLICT_MAX_LINE_CHARS == 20000 and ug.CONFLICT_MAX_PAIR_CHARS == 6000000
    long_line = {"name": "n", "text": "x" * 20001}
    p = ug.check_conflict_documents([long_line, {"name": "m", "text": "short"}])
    assert p["code"] == "TOO_LARGE" and "20,000 characters without a line break" in p["message"] and "PDF or Word" in p["message"]
    assert ug.check_conflict_documents([{"name": "n", "text": "x" * 20000}, {"name": "m", "text": "short"}]) is None


def test_rule_order_is_count_then_characters_then_line_then_pairs():
    assert ug.check_conflict_documents([{"name": "n", "text": "x" * 600001}] * 2)["message"].startswith("One of the documents has more than 600,000")
    assert "long stretch" in ug.check_conflict_documents([{"name": "n", "text": "x" * 30000}])["message"]
    assert "too much text" in ug.check_conflict_documents(_docs(10, 600000))["message"]


def test_json_document_count_cap():
    assert ug.check_json_documents([{}] * 10) is None
    p = ug.check_json_documents([{}] * 11)
    assert p["code"] == "TOO_MANY_FILES" and p["status"] == 413
    assert ug.check_json_documents(None) is None and ug.check_json_documents("text") is None and ug.check_json_documents({"a": 1}) is None


# ── readable-document minimum (conflict routes) ─────────────────────────────────────────────────
def test_readable_count_rule():
    assert ug.check_readable_count(2, 0, 2) is None
    assert ug.check_readable_count(3, 1, 2) is None                      # two readable, one not
    p = ug.check_readable_count(2, 1, 2)
    assert p["code"] == "NOT_ENOUGH_READABLE_TEXT" and p["status"] == 400
    assert p["message"] == ("At least 2 documents with readable text are needed. 1 of your 2 documents had no readable text (for example a scan). "
                            "Upload text versions or run OCR on the scans first.")
    for n, u in ((1, 1), (2, 2), (5, 5)):                                # zero readable is always refused
        p = ug.check_readable_count(n, u, 2)
        assert p["code"] == "NOT_ENOUGH_READABLE_TEXT" and "none of your %d document" % n in p["message"], p
    assert ug.check_readable_count(1, 1, 0)["code"] == "NOT_ENOUGH_READABLE_TEXT"      # a minimum of 0 still needs one readable document
    assert ug.check_readable_count(1, 1, 1)["code"] == "NOT_ENOUGH_READABLE_TEXT"
    assert ug.check_readable_count(2, 1, 1) is None and ug.check_readable_count(1, 0, 0) is None
    assert ug.check_readable_count(3, 2, 2)["message"].count("2 of your 3") == 1


# ── content-type based part rules (pictures are not XML; a name cannot hide an XML bomb) ─────────
STYLES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"
DOC_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
_BODY = "<w:p><w:r><w:t>Hello contract text. Payment within 30 days.</w:t></w:r></w:p>"


def docx_typed(parts=None, overrides=None, defaults=None, content_types=True, link=()):
    """A valid minimal .docx plus extra parts. parts: name -> bytes (stored, no compression, unless the name is in `deflate`)
    or (bytes, "deflate"). overrides: part name -> content type; defaults: extension -> content type. link: names of extra parts
    to relate to the main document, so python-docx really loads them."""
    import zipfile as zf_mod
    ct = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
          '<Default Extension="xml" ContentType="application/xml"/>',
          f'<Override PartName="/word/document.xml" ContentType="{DOC_CT}"/>']
    ct += [f'<Default Extension="{e}" ContentType="{t}"/>' for e, t in (defaults or {}).items()]
    ct += [f'<Override PartName="/{n}" ContentType="{t}"/>' for n, t in (overrides or {}).items()]
    ct.append("</Types>")
    buf = io.BytesIO()
    with zf_mod.ZipFile(buf, "w", zf_mod.ZIP_DEFLATED, compresslevel=6) as z:
        if content_types:
            z.writestr("[Content_Types].xml", "".join(ct))
        z.writestr("_rels/.rels", RELS)
        z.writestr("word/document.xml", doc_xml(_BODY, 1))
        if link:
            rels = "".join(f'<Relationship Id="rIdX{k}" Type="{IMAGE_REL}" Target="{n[len("word/"):]}"/>' for k, n in enumerate(link))
            z.writestr("word/_rels/document.xml.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + rels + "</Relationships>")
        for name, blob in (parts or {}).items():
            if isinstance(blob, tuple):
                z.writestr(name, blob[0], compress_type=zf_mod.ZIP_DEFLATED)
            else:
                z.writestr(name, blob, compress_type=zf_mod.ZIP_STORED)
    return buf.getvalue()


def random_xml(megabytes):
    import base64
    import os
    body = base64.b64encode(os.urandom(int(megabytes * 1024 * 1024 * 3 / 4))).decode()
    return f'<?xml version="1.0"?><w:styles xmlns:w="{W}"><!--{body}--></w:styles>'.encode()


def test_big_incompressible_picture_is_accepted():
    import os
    data = docx_typed({"word/media/photo.png": os.urandom(10 * 1024 * 1024)}, defaults={"png": "image/png"}, link=["word/media/photo.png"])
    assert len(data) > 10 * 1024 * 1024
    assert ug.check_upload_strict(data, "stamp.docx") == ("docx", None)


def test_picture_of_28_megabytes_is_accepted():
    import os
    data = docx_typed({"word/media/photo.png": os.urandom(28 * 1024 * 1024)}, defaults={"png": "image/png"}, link=["word/media/photo.png"])
    assert len(data) < ug.STRICT_MAX_BYTES
    assert ug.check_upload_strict(data, "scan.docx") == ("docx", None)


def test_xml_part_of_9_megabytes_is_refused():
    data = docx_typed({"word/styles.xml": random_xml(9)}, overrides={"word/styles.xml": STYLES_CT})
    kind, problem = ug.check_upload_strict(data, "a.docx")
    assert kind is None and problem["code"] == "UNSAFE_ARCHIVE" and problem["status"] == 422


def test_a_part_named_like_a_picture_but_declared_as_xml_is_refused():
    """python-docx decides by [Content_Types].xml, not by the name, so word/media/x.png can be a 9 MB XML part."""
    data = docx_typed({"word/media/x.png": random_xml(9)}, overrides={"word/media/x.png": STYLES_CT}, defaults={"png": "image/png"})
    assert code_of(ug.check_upload_strict(data, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_a_part_with_no_content_type_is_treated_as_xml():
    data = docx_typed({"word/custom/x.dat": random_xml(9)})
    assert code_of(ug.check_upload_strict(data, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_picture_with_a_ratio_above_200_is_refused():
    data = docx_typed({"word/media/z.png": (b"\x00" * (6 * 1024 * 1024),)}, defaults={"png": "image/png"})
    assert code_of(ug.check_upload_strict(data, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_missing_or_bad_content_types_is_refused():
    assert code_of(ug.check_upload_strict(docx_typed(content_types=False), "a.docx")[1]) == "NOT_A_DOCX"
    import zipfile as zf_mod
    def with_ct(ct_xml):
        buf = io.BytesIO()
        with zf_mod.ZipFile(buf, "w", zf_mod.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", ct_xml)
            z.writestr("_rels/.rels", RELS)
            z.writestr("word/document.xml", doc_xml(_BODY, 1))
        return buf.getvalue()
    assert code_of(ug.check_upload_strict(with_ct("<Types"), "a.docx")[1]) == "NOT_A_DOCX"
    assert code_of(ug.check_upload_strict(with_ct('<!DOCTYPE Types [<!ENTITY a "b">]><Types xmlns="x"/>'), "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_overall_and_xml_totals_are_separate(monkeypatch):
    import os
    monkeypatch.setattr(ug, "STRICT_DOCX_MAX_TOTAL", 100 * 1024)          # XML parts together
    pics = docx_typed({"word/media/a.png": os.urandom(300 * 1024), "word/media/b.png": os.urandom(300 * 1024)}, defaults={"png": "image/png"})
    assert ug.check_upload_strict(pics, "a.docx") == ("docx", None)       # pictures do not count towards the XML total
    xmls = docx_typed({"word/s1.xml": random_xml(0.06), "word/s2.xml": random_xml(0.06)}, overrides={"word/s1.xml": STYLES_CT, "word/s2.xml": STYLES_CT})
    assert code_of(ug.check_upload_strict(xmls, "a.docx")[1]) == "UNSAFE_ARCHIVE"
    monkeypatch.setattr(ug, "STRICT_DOCX_MAX_ALL", 500 * 1024)            # every part together
    assert code_of(ug.check_upload_strict(pics, "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_content_type_classification():
    t = ({"/word/styles.xml": STYLES_CT, "/word/media/x.png": STYLES_CT}, {"png": "image/png", "xml": "application/xml", "bin": "application/octet-stream",
                                                                           "rels": "application/vnd.openxmlformats-package.relationships+xml", "txt": "text/xml; charset=utf-8"})
    assert ug._is_xml_part("word/styles.xml", t) and ug._is_xml_part("word/document.xml", t) and ug._is_xml_part("word/_rels/document.xml.rels", t)
    assert ug._is_xml_part("word/media/x.png", t)                 # an Override by part name beats the extension
    assert not ug._is_xml_part("word/media/y.png", t) and not ug._is_xml_part("word/embeddings/o.bin", t)
    assert ug._is_xml_part("word/notes.txt", t)                   # text/xml with a parameter
    assert ug._is_xml_part("word/unknown.dat", t) and ug._is_xml_part("word/noextension", t)   # unknown type: treated as XML (strict)
    assert ug.STRICT_DOCX_MAX_ALL == 64 * 1024 * 1024


# ── total element count (memory): every element is counted, memory stays flat, the counted-element result is unchanged ─────────
def _tiny_elements_xml(n, seed=7):
    """n tiny elements that are NOT counted (no p, r or tc), in random order so the part still compresses well under a ratio of 200."""
    import random
    names = random.Random(seed).choices("abcdefgh", k=n)
    return (f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>' + "".join(f"<w:{c}/>" for c in names) + "</w:body></w:document>").encode()


def test_total_element_cap_default_and_margin():
    assert ug.STRICT_DOCX_MAX_TOTAL_ELEMENTS == 250000
    assert ug.STRICT_DOCX_MAX_ELEMENTS == 25000                               # the counted-element cap is unchanged


@pytest.mark.parametrize("path", DOCX_FIXTURES, ids=lambda p: p.name)
def test_total_pass_gives_the_same_counted_result_on_real_fixtures(path):
    xml = zipfile.ZipFile(path).read("word/document.xml")
    assert ug.count_docx_elements(xml, 25000, 250000) == ug.count_docx_elements(xml, 25000)


def test_total_pass_counts_every_element_and_stops_early():
    xml = _tiny_elements_xml(3000)
    assert ug.count_docx_elements(xml, 25000, 3000 + 2) == (0, "ok")           # w:document + w:body + 3,000 = 3,002
    assert ug.count_docx_elements(xml, 25000, 3001) == (0, "total")
    t = time.time()
    big = _tiny_elements_xml(400000)
    assert ug.count_docx_elements(big, 25000, 1000) == (0, "total") and time.time() - t < 3
    assert ug.count_docx_elements(big, 25000) == (0, "ok")                    # without total_limit nothing else is counted: old behaviour


def test_counted_cap_still_comes_first_and_keeps_its_numbers():
    xml = doc_xml(para("x"), 15000, prefix="x").encode()                      # 30,000 counted elements
    assert ug.count_docx_elements(xml, 25000, 250000) == (25001, "limit")
    assert ug.count_docx_elements(xml, 25000) == (25001, "limit")


def test_total_pass_still_finds_doctype_and_malformed_xml():
    bad = doc_xml(para(), 3, before_root='<!DOCTYPE w:document [<!ENTITY a "b">]>').encode()
    assert ug.count_docx_elements(bad, 25000, 250000)[1] == "dtd"
    assert ug.count_docx_elements(doc_xml(para(), 3).encode()[:-30], 25000, 250000)[1] == "malformed"
    only_root = f'<?xml version="1.0"?><!DOCTYPE w:document [<!ENTITY a "b">]><w:document xmlns:w="{W}"/>'.encode()
    assert ug.count_docx_elements(only_root, 25000, 250000)[1] == "dtd"


def test_docx_with_a_million_tiny_uncounted_elements_is_refused_fast():
    data = zip_docx(_tiny_elements_xml(1390000).decode(), level=1)                      # about 8 MB of xml, ratio about 10: passes the size and ratio rules
    t = time.time()
    kind, problem = ug.check_upload_strict(data, "tiny.docx")
    print(f"\nTOTAL-COUNT refusal of 1.39M tiny elements took {time.time() - t:.2f}s")
    assert kind is None and problem["code"] == "UNSAFE_ARCHIVE" and "250,000 elements" in problem["message"] and time.time() - t < 4


def test_docx_just_under_the_total_cap_is_accepted():
    data = zip_docx(_tiny_elements_xml(249000).decode(), level=1)                       # 249,000 + 2 elements
    assert ug.check_upload_strict(data, "big.docx") == ("docx", None)
    data = zip_docx(_tiny_elements_xml(250000).decode(), level=1)                       # 250,002 elements: over
    assert code_of(ug.check_upload_strict(data, "big.docx")[1]) == "UNSAFE_ARCHIVE"


# ── ONE element budget for ALL XML parts (python-docx parses styles, numbering, settings, headers, footers ... too) ────────────
RT_BASE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
CT_BASE = "application/vnd.openxmlformats-officedocument.wordprocessingml."
REL_KINDS = {  # part name -> (root element, content type, relationship type)
    "word/styles.xml": ("styles", CT_BASE + "styles+xml", RT_BASE + "styles"),
    "word/numbering.xml": ("numbering", CT_BASE + "numbering+xml", RT_BASE + "numbering"),
    "word/settings.xml": ("settings", CT_BASE + "settings+xml", RT_BASE + "settings"),
    "word/header1.xml": ("hdr", CT_BASE + "header+xml", RT_BASE + "header"),
    "word/footnotes.xml": ("footnotes", CT_BASE + "footnotes+xml", RT_BASE + "footnotes"),
}


def tiny_part(root, n, seed=3):
    """A part of n tiny elements, none of them a Word paragraph, run or cell, in random order so it compresses to a ratio far below 200."""
    import random
    names = random.Random(seed).choices("abcdefgh", k=n)
    return (f'<?xml version="1.0"?><w:{root} xmlns:w="{W}">' + "".join(f"<w:{c}/>" for c in names) + f"</w:{root}>").encode()


def linked_docx(parts, link=True, main_xml=None, rels_type=None):
    """A .docx whose extra XML parts have the right content types and (link=True) are related to the main part, exactly like Word's own
    styles, numbering, settings and header parts; with link=False they are orphans. rels_type: content type declared for the .rels extension."""
    ct = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
          f'<Default Extension="rels" ContentType="{rels_type or "application/vnd.openxmlformats-package.relationships+xml"}"/>',
          '<Default Extension="xml" ContentType="application/xml"/>', f'<Override PartName="/word/document.xml" ContentType="{DOC_CT}"/>']
    ct += [f'<Override PartName="/{n}" ContentType="{REL_KINDS[n][1]}"/>' for n in parts]
    ct.append("</Types>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as z:
        z.writestr("[Content_Types].xml", "".join(ct))
        z.writestr("_rels/.rels", RELS)
        z.writestr("word/document.xml", main_xml or doc_xml(_BODY, 1))
        if link:
            rels = "".join(f'<Relationship Id="rIdL{k}" Type="{REL_KINDS[n][2]}" Target="{n[5:]}"/>' for k, n in enumerate(parts))
            z.writestr("word/_rels/document.xml.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + rels + "</Relationships>")
        for n, blob in parts.items():
            z.writestr(n, blob)
    return buf.getvalue()


def sum_of_xml_elements(data):
    """Independent count: every element of every XML part of the file (by content type), the number the shared budget is spent on."""
    from lxml import etree
    zf = zipfile.ZipFile(io.BytesIO(data))
    types, _ = ug._read_content_types(zf)
    total = 0
    for i in zf.infolist():
        if not i.filename.endswith("/") and ug._is_xml_part(i.filename, types):
            total += sum(1 for _ev, _el in etree.iterparse(io.BytesIO(zf.read(i)), events=("start",)))
    return total


S_XML, N_XML, T_XML, F_XML = "word/styles.xml", "word/numbering.xml", "word/settings.xml", "word/footnotes.xml"


def test_budget_is_a_sum_over_all_xml_parts_and_keeps_its_value():
    assert ug.STRICT_DOCX_MAX_TOTAL_ELEMENTS == 250000 and ug.STRICT_DOCX_MAX_ELEMENTS == 25000


def test_count_xml_elements_counts_everything_and_stops_early():
    xml = tiny_part("styles", 3000)
    assert ug.count_xml_elements(xml, 10 ** 6) == (3001, "ok")                          # the root and 3,000 children
    assert ug.count_xml_elements(xml, 3001) == (3001, "ok")
    assert ug.count_xml_elements(xml, 3000) == (3001, "total")
    assert ug.count_xml_elements(tiny_part("styles", 400000), 1000) == (1001, "total")   # stops at once, does not read on
    assert ug.count_xml_elements(b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY b "c">]><a/>', 100)[1] == "dtd"
    assert ug.count_xml_elements(tiny_part("styles", 50)[:-9], 1000)[1] == "malformed"


def test_bomb_in_styles_xml_is_refused_and_python_docx_is_never_asked(monkeypatch):
    import docx
    monkeypatch.setattr(docx, "Document", lambda *a, **k: (_ for _ in ()).throw(AssertionError("python-docx must not run on a bomb")))
    data = linked_docx({S_XML: tiny_part("styles", 300000)})                              # main part is tiny and fine
    assert zipfile.ZipFile(io.BytesIO(data)).getinfo(S_XML).file_size < ug.STRICT_DOCX_MAX_PART
    t = time.time()
    kind, problem = ug.check_upload_strict(data, "bomb.docx")
    assert kind is None and problem["code"] == "UNSAFE_ARCHIVE" and "250,000 elements" in problem["message"] and time.time() - t < 3


def test_same_bomb_split_over_three_linked_parts_is_refused():
    one = lambda name: ug.check_upload_strict(linked_docx({name: tiny_part(REL_KINDS[name][0], 90000)}), "a.docx")
    assert one(S_XML) == ("docx", None) and one(N_XML) == ("docx", None) and one(T_XML) == ("docx", None)       # each part alone is fine
    data = linked_docx({S_XML: tiny_part("styles", 90000), N_XML: tiny_part("numbering", 90000, 4), T_XML: tiny_part("settings", 90000, 5)})
    assert sum(zipfile.ZipFile(io.BytesIO(data)).getinfo(n).file_size for n in (S_XML, N_XML, T_XML)) < ug.STRICT_DOCX_MAX_TOTAL
    kind, problem = ug.check_upload_strict(data, "split.docx")
    assert kind is None and problem["code"] == "UNSAFE_ARCHIVE" and "250,000 elements" in problem["message"]


def test_the_main_part_and_the_extra_parts_share_the_same_budget():
    main = tiny_part("document", 150000).decode()
    assert ug.check_upload_strict(linked_docx({}, main_xml=main), "a.docx") == ("docx", None)
    assert ug.check_upload_strict(linked_docx({S_XML: tiny_part("styles", 120000)}), "a.docx") == ("docx", None)
    kind, problem = ug.check_upload_strict(linked_docx({S_XML: tiny_part("styles", 120000)}, main_xml=main), "a.docx")
    assert kind is None and problem["code"] == "UNSAFE_ARCHIVE"


def test_parts_that_are_not_linked_or_not_parsed_are_counted_too():
    assert code_of(ug.check_upload_strict(linked_docx({S_XML: tiny_part("styles", 300000)}, link=False), "orphan.docx")[1]) == "UNSAFE_ARCHIVE"
    assert code_of(ug.check_upload_strict(linked_docx({F_XML: tiny_part("footnotes", 300000)}), "notes.docx")[1]) == "UNSAFE_ARCHIVE"


def test_relationship_parts_are_xml_whatever_the_content_type_says():
    assert ug._is_xml_part("word/_rels/document.xml.rels", ({}, {"rels": "application/octet-stream"})) is True
    assert ug._is_xml_part("word/media/photo.png", ({}, {"png": "image/png"})) is False       # pictures are still not XML
    import random
    rnd = random.Random(1)
    rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + "".join(f'<a i="{rnd.randrange(100000)}"/>' for _ in range(300000)) + "</Relationships>").encode()   # varied: the ratio rule must not be what refuses it
    data = linked_docx({S_XML: tiny_part("styles", 10)}, rels_type="application/octet-stream")
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as zout:
        for i in zin.infolist():
            zout.writestr(i.filename, rels if i.filename == "word/_rels/document.xml.rels" else zin.read(i))
    assert code_of(ug.check_upload_strict(buf.getvalue(), "rels.docx")[1]) == "UNSAFE_ARCHIVE"


def test_doctype_in_an_extra_part_is_refused_but_a_broken_extra_part_is_left_to_the_extractor():
    dtd = b'<?xml version="1.0"?><!DOCTYPE w:styles [<!ENTITY a "b">]><w:styles xmlns:w="' + W.encode() + b'"/>'
    assert code_of(ug.check_upload_strict(linked_docx({S_XML: dtd}), "a.docx")[1]) == "UNSAFE_ARCHIVE"
    assert ug.check_upload_strict(linked_docx({S_XML: b"not xml at all"}), "a.docx") == ("docx", None)       # as before this change


def test_a_broken_extra_part_cannot_hide_a_bomb_in_front_of_the_break():
    broken = tiny_part("styles", 300000)[:-40]                         # 300,000 elements, then the file ends in the middle of a tag
    assert code_of(ug.check_upload_strict(linked_docx({S_XML: broken}), "a.docx")[1]) == "UNSAFE_ARCHIVE"


def test_default_python_docx_parts_and_real_fixtures_stay_far_below_the_budget():
    for path in DOCX_FIXTURES:
        data = path.read_bytes()
        total = sum_of_xml_elements(data)
        assert total * 4 <= ug.STRICT_DOCX_MAX_TOTAL_ELEMENTS, (path.name, total)
        assert ug.check_upload_strict(data, path.name) == ("docx", None), path.name


def test_realistic_300_page_contract_sum_over_all_xml_parts_has_a_4x_margin():
    data = realistic_300_docx()
    total = sum_of_xml_elements(data)
    main = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    print(f"\nREALISTIC 300-page docx: SUM over all XML parts={total} (main part counted={ug.count_docx_elements(main, 10 ** 9)[0]}), budget {ug.STRICT_DOCX_MAX_TOTAL_ELEMENTS}")
    assert total * 4 <= ug.STRICT_DOCX_MAX_TOTAL_ELEMENTS
    assert ug.check_upload_strict(data, "contract.docx") == ("docx", None)


def realistic_300_docx():
    import docx
    d = docx.Document()
    for page in range(300):
        if page % 5 == 4:
            t = d.add_table(rows=8, cols=4)
            for row in t.rows:
                for cell in row.cells:
                    cell.paragraphs[0].add_run("Schedule item ")
                    cell.paragraphs[0].add_run("Rs. 4,75,000").bold = True
            for _ in range(3):
                p = d.add_paragraph()
                for r in range(3):
                    p.add_run(f"Payment shall be made within thirty days of invoice {r}. ")
        else:
            for k in range(8 + page % 2):
                p = d.add_paragraph()
                p.add_run(f"{k + 1}. ")
                p.add_run("The Company").bold = True
                p.add_run(" shall indemnify the Contractor against all losses arising from breach of this Agreement. ")
                if k % 2:
                    p.add_run("Disputes are subject to the courts at Mumbai.")
        d.add_page_break()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
