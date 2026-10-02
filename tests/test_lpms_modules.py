"""Practice (LPMS): clients, RTI tracker, calendar, communication, reports (PDF + Excel), search, backup, audit log, dashboard."""
import io
import json
import zipfile

from test_lpms_flow import iso, make_case, today


# ── clients ─────────────────────────────────────────────────────────────────────────
def test_clients_crud_duplicates_privacy(firm):
    t = firm.team
    a, meena, ravi = t["asha"], t["meena"], t["ravi"]
    c = meena.ok(meena.post("/clients", {"name": "Suresh Menon", "phone": "+91 98765 43210", "email": "Suresh@Menon.in", "address": "12 MG Road, Kochi", "id_type": "Aadhaar",
                                         "id_number": "123412341234", "occupation": "Teacher", "comm_pref": "whatsapp"}), 201)["client"]
    assert c["email"] == "suresh@menon.in" and c["id_number"] == "123412341234"
    lst = a.ok(a.get("/clients"))["clients"]
    assert lst[0]["id_number"].endswith("1234") and "•" in lst[0]["id_number"]               # the list never shows the whole ID
    d = ravi.post("/clients", {"name": "Suresh M", "phone": "09876543210"})                  # same number, typed differently
    assert d.status_code == 409 and d.get_json()["code"] == "DUPLICATE_CLIENT"
    ravi.ok(ravi.post("/clients", {"name": "Suresh M", "phone": "09876543210", "force": True}), 201)
    assert a.post("/clients", {"name": ""}).status_code == 400
    assert a.post("/clients", {"name": "Z", "email": "nope"}).status_code == 400
    assert a.post("/clients", {"name": "Z", "phone": "abc"}).status_code == 400
    assert a.post("/clients", {"name": "Z", "comm_pref": "carrier pigeon"}).status_code == 400
    a.ok(a.patch(f"/clients/{c['id']}", {"occupation": "Principal"}))
    assert a.ok(a.get(f"/clients/{c['id']}"))["client"]["occupation"] == "Principal"
    case = make_case(a, client_id=c["id"])
    det = a.ok(a.get(f"/clients/{c['id']}"))["client"]
    assert det["case_list"][0]["id"] == case["id"]
    assert a.ok(a.get("/clients?q=menon"))["total"] == 1 and a.ok(a.get("/clients?q=98765"))["total"] == 2
    a.ok(a.post(f"/clients/{c['id']}/archive", {"archive": True}))
    assert a.ok(a.get("/clients?q=menon"))["total"] == 0 and a.ok(a.get("/clients?archived=1"))["total"] == 1
    # a client of another practice does not exist as far as we are concerned
    other = firm.person("Other Firm Owner")
    other.ok(other.post("/firm", {"name": "Other & Co"}), 201)
    assert other.get(f"/clients/{c['id']}").status_code == 404
    assert other.post("/cases", {"case_no": "X 1", "court": "C", "title": "t", "client_id": c["id"]}).status_code == 400
    assert other.get(f"/cases/{case['id']}").status_code == 404


def test_two_practices_never_see_each_other(firm):
    a = firm.team["asha"]
    case = make_case(a, case_no="ZZ 1/2026", title="Rao private matter", client={"name": "Rao Client"})
    a.ok(a.post(f"/cases/{case['id']}/hearings", {"date": iso(0)}), 201)
    rival = firm.person("Rival Lawyer")
    rival.ok(rival.post("/firm", {"name": "Rival Chambers"}), 201)
    make_case(rival, case_no="ZZ 1/2026", title="Rival matter")                              # same number is fine in another practice
    assert rival.ok(rival.get("/cases"))["total"] == 1
    assert rival.ok(rival.get("/search?q=Rao"))["cases"] == [] and rival.ok(rival.get("/search?q=Rao"))["clients"] == []
    assert rival.ok(rival.get(f"/hearings?date={iso(0)}"))["hearings"] == []
    for url in (f"/cases/{case['id']}", f"/cases/{case['id']}/proceedings", f"/cases/{case['id']}/timeline", f"/cases/{case['id']}/notes"):
        assert rival.get(url).status_code == 404
    assert rival.patch(f"/cases/{case['id']}", {"remarks": "x"}).status_code == 404
    assert rival.ok(rival.get("/members"))["members"][0]["name"] == "Rival Lawyer" and len(rival.ok(rival.get("/members"))["members"]) == 1
    assert rival.ok(rival.get("/audit"))["total"] >= 1 and all("Rao" not in (x["summary"] or "") for x in rival.ok(rival.get("/audit"))["items"])


# ── communication ───────────────────────────────────────────────────────────────────
def test_whatsapp_and_email_reminders(firm):
    a, ravi = firm.team["asha"], firm.team["ravi"]
    c = make_case(a, advocate=ravi.member_id, client={"name": "Anita Roy", "phone": "98765 43210", "email": "anita@roy.in"}, title="Roy v. Roy")["id"]
    assert a.post(f"/cases/{c}/whatsapp", {}).status_code == 409                                  # nothing to remind about yet
    a.ok(a.post(f"/cases/{c}/hearings", {"date": iso(2), "time": "11:00", "hall_no": "3"}), 201)
    d = a.ok(a.get(f"/cases/{c}/message-draft"))
    assert "Anita Roy" in d["body"] and "11:00" in d["body"] and d["client"]["phone"] is True
    w = ravi.ok(ravi.post(f"/cases/{c}/whatsapp", {}))
    assert w["url"].startswith("https://wa.me/919876543210?text=") and "Roy%20v.%20Roy" in w["url"]
    custom = ravi.ok(ravi.post(f"/cases/{c}/whatsapp", {"text": "Please bring your ID"}))
    assert custom["url"].endswith("Please%20bring%20your%20ID")
    hist = a.ok(a.get(f"/comms?case_id={c}"))["comms"]
    assert [x["status"] for x in hist if x["channel"] == "whatsapp"] == ["opened", "opened"]       # honest: opened, not "delivered"
    # e-mail: the server says plainly when it is not set up, and still records the attempt
    r = a.ok(a.post(f"/cases/{c}/email", {}))
    assert r["delivered"] is False and r["configured"] is False and "not set up" in r["message"] and firm.sent == []
    firm.mail_ok = True
    r = a.ok(a.post(f"/cases/{c}/email", {"subject": "About your case", "body": "Hello Anita"}))
    assert r["delivered"] is True and firm.sent[-1]["to"] == "anita@roy.in" and firm.sent[-1]["subject"] == "About your case"
    firm.mail_returns = False
    assert a.ok(a.post(f"/cases/{c}/email", {"body": "again"}))["delivered"] is False
    stat = [x["status"] for x in a.ok(a.get(f"/comms?case_id={c}&channel=email"))["comms"]]
    assert stat == ["failed", "sent", "not_sent"]
    # missing contact details are explained, not swallowed
    c2 = make_case(a, case_no="N 2", client={"name": "No Phone"})["id"]
    a.ok(a.post(f"/cases/{c2}/hearings", {"date": iso(2)}), 201)
    assert a.post(f"/cases/{c2}/whatsapp", {}).json["code"] == "NO_PHONE"
    assert a.post(f"/cases/{c2}/email", {}).json["code"] == "NO_EMAIL"
    c3 = make_case(a, case_no="N 3")["id"]
    assert a.post(f"/cases/{c3}/whatsapp", {}).json["code"] == "NO_CLIENT"
    # switched off in settings -> refused
    a.ok(a.put("/settings", {"channels": {"whatsapp": False}}))
    assert a.post(f"/cases/{c}/whatsapp", {}).status_code == 409


# ── RTI tracker ─────────────────────────────────────────────────────────────────────
def test_rti_tracker_deadlines_and_appeals(firm):
    a, meena = firm.team["asha"], firm.team["meena"]
    case = make_case(a)["id"]
    assert a.post("/rti", {"subject": "Copy of file notings"}).status_code == 400                    # department missing
    assert a.post("/rti", {"subject": "x", "department": "MCD", "status": "filed"}).status_code == 400  # filed needs a date
    r = meena.ok(meena.post("/rti", {"subject": "Copy of file notings", "department": "MCD, Civil Lines Zone", "case_id": case, "filing_date": iso(-20), "status": "filed", "fee": "10"}), 201)["rti"]
    assert r["response_due"] == iso(10) and r["next_deadline"]["label"] == "Reply due" and r["next_deadline"]["days"] == 10     # 30 days from filing
    rid = r["id"]
    # the reply arrives -> the first-appeal window opens (30 days from the reply), and the status change is recorded
    r = a.ok(a.patch(f"/rti/{rid}", {"status": "partial", "response_date": iso(-2)}))["rti"]
    assert r["appeal_due"] == iso(28) and r["next_deadline"]["label"] == "First appeal due"
    a.ok(a.patch(f"/rti/{rid}", {"status": "appeal1", "appeal1_date": iso(0), "appeal2_due": iso(90)}))
    a.ok(a.post(f"/rti/{rid}/events", {"note": "Hearing before FAA on 15th"}), 201)
    det = a.ok(a.get(f"/rti/{rid}"))
    kinds = [e["kind"] for e in det["events"]]
    assert kinds[0] == "note" and kinds.count("status") == 2 and "created" in kinds and det["rti"]["next_deadline"]["label"] == "Second appeal due"
    assert a.ok(a.get("/rti?status=open"))["rti"][0]["id"] == rid and a.ok(a.get(f"/rti?case_id={case}"))["counts"] == {"appeal1": 1}
    assert any("RTI" in x["title"] for x in a.ok(a.get(f"/cases/{case}/timeline"))["items"])
    assert meena.delete(f"/rti/{rid}").status_code == 403
    a.ok(a.delete(f"/rti/{rid}"))
    assert a.get(f"/rti/{rid}").status_code == 404


# ── calendar ────────────────────────────────────────────────────────────────────────
def test_calendar_merges_everything_and_holidays(firm):
    t = firm.team
    a, ravi, meena = t["asha"], t["ravi"], t["meena"]
    c = make_case(a, advocate=ravi.member_id, next_action="File written statement", next_action_due=iso(4))["id"]
    a.ok(a.post(f"/cases/{c}/hearings", {"date": iso(2), "time": "10:00"}), 201)
    meena.ok(meena.post("/events", {"kind": "appointment", "title": "Client meeting — Mr Rao", "start_date": iso(3), "start_time": "15:00", "end_time": "16:00", "member_id": ravi.member_id, "case_id": c}), 201)
    meena.ok(meena.post("/events", {"kind": "meeting", "title": "Chamber meeting", "start_date": iso(1)}), 201)
    a.ok(a.post("/events", {"kind": "deadline", "title": "Filing deadline", "start_date": iso(5), "case_id": c}), 201)
    a.ok(a.post("/events", {"kind": "holiday", "title": "Winter vacation", "start_date": iso(6), "end_date": iso(8)}), 201)
    a.ok(a.post("/rti", {"subject": "Status of road repair", "department": "PWD", "filing_date": iso(-25), "status": "filed"}), 201)
    a.ok(a.post("/followups", {"case_id": c, "due_date": iso(1), "note": "Call client"}), 201)
    cal = a.ok(a.get(f"/calendar?from={iso(0)}&to={iso(10)}"))
    kinds = [i["type"] for i in cal["items"]]
    for k in ("hearing", "appointment", "meeting", "deadline", "holiday", "rti", "followup"):
        assert k in kinds, (k, kinds)
    assert sum(1 for i in cal["items"] if i["type"] == "holiday") == 3                               # a vacation spans its days
    assert [i["date"] for i in cal["items"]] == sorted(i["date"] for i in cal["items"])
    assert a.get("/calendar?from=2020-01-01&to=2026-12-31").status_code == 400
    # staff manage appointments but only a senior adds holidays
    assert meena.post("/events", {"kind": "holiday", "title": "Nope", "start_date": iso(2)}).status_code == 403
    ev = next(i for i in cal["items"] if i["type"] == "meeting")["event"]
    assert ravi.patch(f"/events/{ev['id']}", {"title": "x"}).status_code == 403                    # not theirs
    meena.ok(meena.patch(f"/events/{ev['id']}", {"title": "Chamber meeting (moved)"}))
    assert a.post("/events", {"kind": "meeting", "title": "Bad", "start_date": iso(3), "end_date": iso(1)}).status_code == 400
    meena.ok(meena.delete(f"/events/{ev['id']}"))
    # the fixed national holidays button is idempotent
    assert a.ok(a.post("/events/holidays/fixed", {"year": 2027}))["added"] == 5
    assert a.ok(a.post("/events/holidays/fixed", {"year": 2027}))["added"] == 0
    # Ravi was told about his appointment
    assert any("Client meeting" in n["title"] for n in ravi.ok(ravi.get("/notifications"))["items"])


# ── reports ─────────────────────────────────────────────────────────────────────────
def _seed_report_data(firm):
    t = firm.team
    a, ravi = t["asha"], t["ravi"]
    c1 = make_case(a, advocate=ravi.member_id, case_no="R 1/2026", title="=cmd|' /C calc'!A0", case_type="Civil", client={"name": "Report Client"}, category="Property")["id"]
    c2 = make_case(a, advocate=a.member_id, case_no="R 2/2026", title="Closed Matter", case_type="Criminal")["id"]
    a.ok(a.post(f"/cases/{c1}/hearings", {"date": iso(0), "purpose": "Evidence"}), 201)
    a.ok(a.post(f"/cases/{c1}/hearings", {"date": iso(7)}), 201)
    ravi.ok(ravi.post(f"/cases/{c1}/proceedings", {"notes": "PW1 examined", "outcome": "heard"}), 201)
    a.ok(a.patch(f"/cases/{c2}", {"status": "Disposed", "outcome": "Acquitted"}))
    return c1, c2


def test_every_report_builds_in_json_xlsx_and_pdf(firm):
    import openpyxl
    a = firm.team["asha"]
    _seed_report_data(firm)
    month = today().strftime("%Y-%m")
    for kind in ("monthly_summary", "advocate_performance", "case_status", "upcoming_hearings", "closed_cases", "client_activity", "cause_list"):
        rep = a.ok(a.get(f"/reports/{kind}?month={month}"))
        assert rep["title"] and rep["firm"] == "Rao & Associates" and rep["summary"] is not None and rep["tables"]
        x = a.get(f"/reports/{kind}?format=xlsx&month={month}")
        assert x.status_code == 200 and "spreadsheetml" in x.mimetype and x.data[:2] == b"PK"
        wb = openpyxl.load_workbook(io.BytesIO(x.data))
        assert "Summary" in wb.sheetnames
        p = a.get(f"/reports/{kind}?format=pdf&month={month}")
        assert p.status_code == 200 and p.mimetype == "application/pdf" and p.data[:5] == b"%PDF-" and len(p.data) > 2000
    # numbers are right, not just present
    m = {s["label"]: s["value"] for s in a.ok(a.get(f"/reports/monthly_summary?month={month}"))["summary"]}
    assert m["New cases"] == 2 and m["Cases closed"] == 1 and m["Proceedings recorded"] == 1 and m["Hearings held or adjourned"] == 1
    perf = a.ok(a.get("/reports/advocate_performance"))["tables"][0]["rows"]
    ravi = next(r for r in perf if r["advocate"] == "Ravi Nair")
    assert ravi["active"] == 1 and ravi["proceedings"] == 1 and ravi["hearings"] == 1
    closed = a.ok(a.get("/reports/closed_cases"))["tables"][0]["rows"]
    assert [r["case_no"] for r in closed] == ["R 2/2026"] and closed[0]["outcome"] == "Acquitted"
    up = a.ok(a.get("/reports/upcoming_hearings?days=30"))["tables"][0]["rows"]
    assert len(up) == 2 or len(up) == 1                                                             # today's is included while still scheduled
    cl = a.ok(a.get(f"/reports/cause_list?date={iso(0)}"))
    assert cl["summary"][0]["value"] == 1                                                           # the one that is now "heard" still appears on its day
    # a case title that looks like a formula is stored as text in Excel, never as a formula
    wb = openpyxl.load_workbook(io.BytesIO(a.get("/reports/case_status?format=xlsx").data))
    cells = [c.value for row in wb["Cases"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "'=cmd|' /C calc'!A0" in cells and "=cmd|' /C calc'!A0" not in cells
    assert a.get("/reports/nope").status_code == 404 and a.get("/reports/case_status?format=docx").status_code == 400
    assert a.get("/reports/monthly_summary?month=2026-13").status_code == 400


def test_report_access_by_role_and_export_is_audited(firm):
    t = firm.team
    assert t["ravi"].get("/reports/advocate_performance").status_code == 403
    assert t["meena"].get("/reports/monthly_summary").status_code == 403
    assert t["meena"].ok(t["meena"].get("/reports/client_activity"))
    assert [r["key"] for r in t["ravi"].ok(t["ravi"].get("/reports"))["reports"]] == ["monthly_summary", "case_status", "upcoming_hearings", "closed_cases", "cause_list"]
    t["asha"].ok(t["asha"].get("/reports/case_status?format=pdf"))
    assert any(x["action"] == "report_export" for x in t["asha"].ok(t["asha"].get("/audit?group=reports"))["items"])


# ── backup ──────────────────────────────────────────────────────────────────────────
def test_backup_is_a_complete_zip_without_secrets(firm):
    a = firm.team["asha"]
    _seed_report_data(firm)
    assert firm.team["ravi"].get("/backup/export").status_code == 403
    r = a.get("/backup/export")
    assert r.status_code == 200 and r.mimetype == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.data))
    names = set(z.namelist())
    assert {"cases.json", "hearings.json", "proceedings.json", "clients.json", "audit_log.json", "cases.xlsx", "README.txt"} <= names
    cases = json.loads(z.read("cases.json"))
    assert len(cases) == 2 and {c["case_no"] for c in cases} == {"R 1/2026", "R 2/2026"}
    blob = b"".join(z.read(n) for n in names)
    assert b"TempPass123" not in blob and b"password" not in z.read("members.json").lower()


# ── audit log ───────────────────────────────────────────────────────────────────────
def test_audit_log_is_a_verifiable_chain_for_seniors_only(firm):
    t = firm.team
    a = t["asha"]
    c, _ = _seed_report_data(firm)
    assert t["ravi"].get("/audit").status_code == 403 and t["meena"].get("/audit/verify").status_code == 403
    log = a.ok(a.get("/audit?per_page=200"))
    actions = {x["action"] for x in log["items"]}
    assert {"firm_create", "member_add", "case_create", "case_update", "hearing_create", "proceeding_create"} <= actions
    assert a.ok(a.get("/audit/verify"))["ok"] is True
    assert a.ok(a.get("/audit?group=cases"))["total"] >= 3 and all(x["action"].startswith("case_") for x in a.ok(a.get("/audit?group=cases"))["items"])
    assert a.ok(a.get("/audit?q=ravi"))["total"] >= 1
    csv_ = a.get("/audit/export.csv")
    assert csv_.status_code == 200 and b"case_create" in csv_.data
    # changing a row behind the application's back breaks the chain, and the check says where
    firm.sql("UPDATE lpms_audit SET summary = 'Nothing happened here' WHERE id = 2")
    v = a.ok(a.get("/audit/verify"))
    assert v["ok"] is False and v["broken_at"] == 2


def test_dashboard_numbers_and_pending_actions(firm):
    t = firm.team
    a, ravi, kiran = t["asha"], t["ravi"], t["kiran"]
    c1 = make_case(a, advocate=ravi.member_id, case_no="D 1", title="Today Matter", court="Court A", category="Family", case_type="Family")["id"]
    c2 = make_case(a, advocate=kiran.member_id, case_no="D 2", title="Overdue Matter", next_action="Send notice", next_action_due=iso(-2))["id"]
    c3 = make_case(a, advocate=ravi.member_id, case_no="D 3", title="Undated Matter")["id"]
    a.ok(a.post(f"/cases/{c1}/hearings", {"date": iso(0), "hall_no": "2", "serial_no": "5"}), 201)
    a.ok(a.post(f"/cases/{c1}/hearings", {"date": iso(-3)}), 201)                                       # happened, never recorded
    a.ok(a.post(f"/cases/{c3}/hearings", {"date": iso(4)}), 201)
    d = a.ok(a.get("/dashboard"))
    assert d["counts"]["active"] == 3 and d["counts"]["today"] == 1 and d["counts"]["week"] >= 2
    assert d["today"][0]["title"] == "Today Matter" and d["today"][0]["hall_no"] == "2" and d["today"][0]["serial_no"] == "5"
    kinds = {p["kind"] for p in d["pending"]}
    assert {"hearing_update", "action", "no_date"} <= kinds and d["pending"][0]["severity"] == "overdue"
    assert {w["name"]: w["active"] for w in d["workload"]} == {"Asha Rao": 0, "Kiran Shah": 1, "Ravi Nair": 2}
    assert [x["label"] for x in d["summary"]["types"]][0] in ("Civil", "Family")
    assert d["recent"][0]["id"] in (c1, c2, c3)
    # a junior's default view is their own work
    dj = ravi.ok(ravi.get("/dashboard"))
    assert dj["scope"] == "mine" and dj["counts"]["active"] == 2 and all(p["kind"] != "action" for p in dj["pending"])
    assert ravi.ok(ravi.get("/dashboard?scope=firm"))["counts"]["active"] == 3
    # recording what happened clears the pending item
    ravi.ok(ravi.post(f"/cases/{c1}/proceedings", {"notes": "Done", "proc_date": iso(-3), "hearing_id": a.ok(a.get(f"/cases/{c1}"))["case"]["hearings"][1]["id"]}), 201)
    assert not any(p["kind"] == "hearing_update" for p in a.ok(a.get("/dashboard"))["pending"])
    # staff get a dashboard too, with no pending-work nagging
    ms = t["meena"].ok(t["meena"].get("/dashboard"))
    assert ms["scope"] == "mine" and ms["counts"]["active"] == 0 and t["meena"].ok(t["meena"].get("/dashboard?scope=firm"))["counts"]["active"] == 3


def test_report_pdf_has_a_footer_on_every_page_and_no_blank_last_page():
    import pymupdf
    from utils import lpms_reports as R
    rep = {"title": "T", "firm": "Rao & Associates", "period": "Oct 2026", "generated": "02 Oct 2026", "summary": [{"label": "a", "value": 1}],
           "tables": [{"title": "Rows", "columns": [{"key": "a", "label": "A"}, {"key": "b", "label": "B"}], "rows": [{"a": i, "b": "x" * 10} for i in range(150)]}]}
    doc = pymupdf.open(stream=R.to_pdf(rep), filetype="pdf")
    assert len(doc) >= 3
    for i, page in enumerate(doc, 1):
        text = page.get_text()
        assert f"page {i} of {len(doc)}" in text and "Rao & Associates" in text
        assert len(text.strip().splitlines()) > 2                  # no page that holds only the footer
    short = dict(rep, tables=[dict(rep["tables"][0], rows=rep["tables"][0]["rows"][:3])])
    assert len(pymupdf.open(stream=R.to_pdf(short), filetype="pdf")) == 1


def test_reminder_text_uses_readable_dates():
    from utils import lpms_store as L
    t = L.client_reminder_text("Ravi", "Rao & Associates", "A vs B", "OS 1/2026", "2026-10-11", "14:30", "City Civil Court", "4")
    assert "11 Oct 2026 at 2:30 pm" in t and "2026-10-11" not in t
    assert L.nice_date("2026-09-29") == "29 Sep 2026" and L.nice_date("junk") == "junk"
