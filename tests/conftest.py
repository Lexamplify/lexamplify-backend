"""
Shared fixtures for the Document Hub tests.

They mount the REAL blueprint (routes/dms_routes.py) on a small Flask app with real JWT auth and a
throw-away SQLite database, so uploads, extraction, OCR, search and permissions all run for real.
Nothing about the rest of LexAmplify (Postgres users, Groq, Render) is needed.
"""
import io
import os
import sys
import time

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

VAULT_SCHEMA = """
CREATE TABLE IF NOT EXISTS case_vault (id INTEGER PRIMARY KEY AUTOINCREMENT, case_id TEXT, title TEXT, doc_type TEXT, content TEXT,
    created_at DATETIME DEFAULT CURRENT_TIMESTAMP, folder_id INTEGER, smart_title TEXT, tags TEXT, file_blob BLOB, file_format TEXT,
    user_id INTEGER, link_shared BOOLEAN DEFAULT 0);
CREATE TABLE IF NOT EXISTS vault_folders (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, parent_id INTEGER, user_id INTEGER,
    created_at DATETIME DEFAULT CURRENT_TIMESTAMP, protected BOOLEAN DEFAULT 0, link_shared BOOLEAN DEFAULT 0);
CREATE TABLE IF NOT EXISTS document_vault_shares (id INTEGER PRIMARY KEY AUTOINCREMENT, node_type TEXT NOT NULL, node_id INTEGER NOT NULL,
    team_member_id INTEGER NOT NULL, permission TEXT NOT NULL DEFAULT 'view', created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(node_type, node_id, team_member_id));
CREATE TABLE IF NOT EXISTS teams (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, description TEXT, is_private BOOLEAN DEFAULT 0,
    owner_user_id INTEGER NOT NULL, created_at DATETIME DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS team_memberships (id INTEGER PRIMARY KEY AUTOINCREMENT, team_id INTEGER NOT NULL, user_id INTEGER NOT NULL,
    role TEXT NOT NULL DEFAULT 'member', created_at DATETIME DEFAULT CURRENT_TIMESTAMP, UNIQUE(team_id, user_id));
CREATE TABLE IF NOT EXISTS matters (id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'open',
    team_id INTEGER, lead_counsel TEXT, opened_date DATE, owner_user_id INTEGER NOT NULL, created_at DATETIME DEFAULT CURRENT_TIMESTAMP);
"""


class Harness:
    def __init__(self, tmp_path, monkeypatch, workers=2, encryption_key=None):
        import sqlite3
        from flask import Flask
        from flask_jwt_extended import JWTManager, create_access_token
        monkeypatch.setenv("DMS_STORAGE_DIR", str(tmp_path / "files"))
        monkeypatch.setenv("DMS_WORKERS", str(workers))
        monkeypatch.setenv("DMS_STRICT_MOUNT", "1")          # a failure to mount the paper-to-digital tools must fail the test, not be logged
        monkeypatch.setenv("DMS_SCAN_WORKERS", "1")
        if encryption_key:
            monkeypatch.setenv("DMS_ENCRYPTION_KEY", encryption_key)
        else:
            monkeypatch.delenv("DMS_ENCRYPTION_KEY", raising=False)
        self.db_path = str(tmp_path / "lex.db")
        boot = sqlite3.connect(self.db_path)
        boot.executescript(VAULT_SCHEMA)
        boot.commit()
        boot.close()

        from routes import dms_routes
        importlib_reload(dms_routes)
        self.dms_routes = dms_routes

        self.shared_override = None       # tests may plug in app.py's own sharing functions
        self.folder_override = None

        def shared(uid):
            if self.shared_override:
                return self.shared_override(uid)
            c = sqlite3.connect(self.db_path)
            try:
                docs = {r[0] for r in c.execute("SELECT node_id FROM document_vault_shares WHERE node_type='document' AND team_member_id=?", (uid,))}
                folders = {r[0] for r in c.execute("SELECT node_id FROM document_vault_shares WHERE node_type='folder' AND team_member_id=?", (uid,))}
                perms = {r[0]: r[1] for r in c.execute("SELECT node_id, permission FROM document_vault_shares WHERE node_type='document' AND team_member_id=?", (uid,))}
            finally:
                c.close()
            return folders, docs, lambda doc_id: perms.get(doc_id)

        def folder_access(folder_id, uid, require_edit=True):
            if self.folder_override:
                return self.folder_override(folder_id, uid, require_edit)
            c = sqlite3.connect(self.db_path)
            try:
                r = c.execute("SELECT user_id FROM vault_folders WHERE id=?", (folder_id,)).fetchone()
                if not r:
                    return False
                if r[0] is None or int(r[0]) == int(uid):
                    return True
                s = c.execute("SELECT permission FROM document_vault_shares WHERE node_type='folder' AND node_id=? AND team_member_id=?", (folder_id, uid)).fetchone()
                return bool(s) and (s[0] == "edit" or not require_edit)
            finally:
                c.close()

        self.app = Flask(__name__)
        self.app.config.update(JWT_SECRET_KEY="test-secret-key-that-is-long-enough-for-hs256", JWT_TOKEN_LOCATION=["headers"], TESTING=True,
                               MAX_CONTENT_LENGTH=100 * 1024 * 1024)
        JWTManager(self.app)
        self.bp = dms_routes.create_dms_blueprint({"db_path": self.db_path, "shared": shared, "folder_access": folder_access})
        self.app.register_blueprint(self.bp)
        def _tok(uid):
            with self.app.app_context():
                return create_access_token(identity=str(uid))
        self._tok = _tok
        self.tmp = tmp_path

    def client(self, uid):
        return Client(self, uid)

    def sql(self, q, *a):
        import sqlite3
        c = sqlite3.connect(self.db_path)
        c.row_factory = sqlite3.Row
        try:
            r = c.execute(q, a).fetchall()
            c.commit()
            return r
        finally:
            c.close()

    def stop(self):
        try:
            self.bp.worker.stop(timeout=10)
        except Exception:
            pass


def importlib_reload(mod):
    import importlib
    importlib.reload(mod)


class Client:
    def __init__(self, h, uid):
        self.h, self.uid = h, uid
        self.c = h.app.test_client()
        self.headers = {"Authorization": f"Bearer {h._tok(uid)}"}

    def get(self, url, **kw):
        return self.c.get(url, headers=self.headers, **kw)

    def post(self, url, **kw):
        return self.c.post(url, headers=self.headers, **kw)

    def patch(self, url, json=None):
        return self.c.patch(url, headers=self.headers, json=json)

    def delete(self, url):
        return self.c.delete(url, headers=self.headers)

    def upload(self, name, data, **fields):
        fields = {k: str(v) for k, v in fields.items() if v is not None}
        return self.post("/api/dms/upload", data={"file": (io.BytesIO(data), name), **fields}, content_type="multipart/form-data")

    def new_version(self, doc_id, name, data, note=None):
        return self.post(f"/api/dms/docs/{doc_id}/versions", data={"file": (io.BytesIO(data), name), **({"note": note} if note else {})},
                         content_type="multipart/form-data")

    def wait(self, ids, timeout=90):
        ids = [ids] if isinstance(ids, int) else list(ids)
        t0 = time.time()
        while time.time() - t0 < timeout:
            pending = [i for i in ids if self.get(f"/api/dms/docs/{i}").get_json()["doc"]["status"] in ("queued", "processing")]
            if not pending:
                return
            time.sleep(0.15)
        raise AssertionError(f"documents still processing after {timeout}s: {pending}")

    def docs(self, **params):
        qs = "&".join(f"{k}={v}" for k, v in params.items())
        r = self.get("/api/dms/docs?" + qs)
        assert r.status_code == 200, r.get_json()
        return r.get_json()

    def ids(self, **params):
        return [d["id"] for d in self.docs(**params)["docs"]]


@pytest.fixture
def hub(tmp_path, monkeypatch):
    h = Harness(tmp_path, monkeypatch)
    yield h
    h.stop()


@pytest.fixture
def hub_factory(tmp_path, monkeypatch):
    made = []

    def make(**kw):
        h = Harness(tmp_path, monkeypatch, **kw)
        made.append(h)
        return h
    yield make
    for h in made:
        h.stop()


# ── sample file builders ─────────────────────────────────────────────────────────────
ORDER_TEXT = """IN THE HIGH COURT OF DELHI AT NEW DELHI
W.P.(C) 1234/2024
Rajesh Kumar ... Petitioner
versus
Union of India & Ors. ... Respondents
ORDER
Date: 12-03-2025
Heard learned counsel for the parties. Issue notice. List on 15.05.2025 for further hearing.
The respondents shall file counter affidavit within four weeks."""


def make_pdf(pages):
    import fitz
    doc = fitz.open()
    for t in pages:
        pg = doc.new_page()
        pg.insert_textbox(fitz.Rect(56, 56, 540, 780), t, fontsize=11)
    b = doc.tobytes()
    doc.close()
    return b


def make_scanned_pdf(text_pages, dpi=170):
    """A PDF whose pages are pictures of text only - what a flatbed scanner produces."""
    import fitz
    src = fitz.open()
    for t in text_pages:
        pg = src.new_page()
        pg.insert_textbox(fitz.Rect(56, 56, 540, 780), t, fontsize=13)
    out = fitz.open()
    for pg in src:
        pix = pg.get_pixmap(dpi=dpi)
        p = out.new_page(width=pg.rect.width, height=pg.rect.height)
        p.insert_image(p.rect, stream=pix.tobytes("png"))
    b = out.tobytes()
    src.close()
    out.close()
    return b


def make_broken_font_pdf(real_text):
    """A page whose PICTURE shows real_text but whose hidden text layer is garbage - the failure
    mode of legacy Indian-language PDFs. Extraction must fall back to OCR."""
    import fitz
    src = fitz.open()
    pg = src.new_page()
    pg.insert_textbox(fitz.Rect(56, 56, 540, 780), real_text, fontsize=13)
    pix = pg.get_pixmap(dpi=170)
    out = fitz.open()
    p = out.new_page(width=pg.rect.width, height=pg.rect.height)
    p.insert_image(p.rect, stream=pix.tobytes("png"))
    p.insert_text((60, 100), "(cid:12)(cid:45)(cid:99)(cid:7) " * 12, fontsize=9, render_mode=3)
    b = out.tobytes()
    return b


def make_docx(paragraphs, table=None):
    import docx
    d = docx.Document()
    for t in paragraphs:
        d.add_paragraph(t)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                t.cell(i, j).text = str(v)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(rows, title="Sheet1"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_png_text(lines, size=(1300, 420)):
    from PIL import Image, ImageDraw, ImageFont
    im = Image.new("RGB", size, "white")
    dr = ImageDraw.Draw(im)
    try:
        f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 40)
    except Exception:
        f = ImageFont.load_default()
    for i, ln in enumerate(lines):
        dr.text((40, 40 + i * 80), ln, font=f, fill="black")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def ocr_ok():
    from utils import dms_extract as X
    return bool(X.ocr_engine().get("available"))


# ── Practice (Legal Practice Management) ────────────────────────────────────────────
class PracticeHarness:
    """The REAL Practice blueprint (and the Document Hub next to it, so case documents can be tested) on a small Flask app with
    real JWT auth, a throw-away SQLite file, and an in-memory stand-in for the users database."""

    def __init__(self, tmp_path, monkeypatch, email_configured=False):
        import sqlite3
        from flask import Flask
        from flask_jwt_extended import JWTManager, create_access_token
        monkeypatch.setenv("DMS_STORAGE_DIR", str(tmp_path / "files"))
        monkeypatch.setenv("DMS_WORKERS", "1")
        monkeypatch.setenv("DMS_STRICT_MOUNT", "1")
        monkeypatch.setenv("DMS_SCAN_WORKERS", "1")
        monkeypatch.setenv("LPMS_SCHEDULER", "0")
        self.db_path = str(tmp_path / "lex.db")
        boot = sqlite3.connect(self.db_path)
        boot.executescript(VAULT_SCHEMA)
        boot.commit()
        boot.close()

        self.users, self.passwords, self.sent = {}, {}, []
        self.mail_ok = email_configured
        self.mail_returns = True

        def get_user(uid):
            return self.users.get(int(uid))

        def find_user(email):
            return next((u for u in self.users.values() if u["email"] == email.lower()), None)

        def create_user(name, email, password, phone=None):
            return self.add_user(name, email, password, phone)["id"]

        def set_password(uid, pw):
            self.passwords[int(uid)] = pw

        def send_email(to, subject, body):
            self.sent.append({"to": to, "subject": subject, "body": body})
            return self.mail_returns

        from routes import dms_routes, lpms_routes
        import importlib
        importlib.reload(lpms_routes)
        from utils import lpms_store
        self.store = lpms_store
        self.app = Flask(__name__)
        self.app.config.update(JWT_SECRET_KEY="test-secret-key-that-is-long-enough-for-hs256", JWT_TOKEN_LOCATION=["headers"], TESTING=True,
                               MAX_CONTENT_LENGTH=100 * 1024 * 1024)
        JWTManager(self.app)
        self.lpms = lpms_routes.create_lpms_blueprint({
            "db_path": self.db_path, "get_user": get_user, "find_user": find_user, "create_user": create_user, "set_password": set_password,
            "send_email": send_email, "email_configured": lambda: self.mail_ok, "sync_email": True, "no_scheduler": True})
        self.app.register_blueprint(self.lpms)
        self.dms = dms_routes.create_dms_blueprint({
            "db_path": self.db_path, "shared": lambda uid: (set(), set(), lambda doc_id: None),
            "folder_access": lambda folder_id, uid, require_edit=True: True,
            "on_lpms_document": lambda case_id, doc_id, title, uid, action: lpms_store.record_document_event(self.db_path, case_id, doc_id, title, uid, action)})
        self.app.register_blueprint(self.dms)

        def _tok(uid):
            with self.app.app_context():
                return create_access_token(identity=str(uid))
        self._tok = _tok
        self.tmp = tmp_path

    def add_user(self, name, email, password="TestPass123!", phone=None):
        uid = len(self.users) + 1
        self.users[uid] = {"id": uid, "name": name, "email": email.lower(), "phone": phone}
        self.passwords[uid] = password
        return self.users[uid]

    def person(self, name, email=None, ip="10.1.1.1"):
        u = self.add_user(name, email or f"{name.lower().replace(' ', '.')}@example.com")
        return PracticeClient(self, u["id"], ip)

    def sql(self, q, *a):
        import sqlite3
        c = sqlite3.connect(self.db_path)
        c.row_factory = sqlite3.Row
        try:
            r = c.execute(q, a).fetchall()
            c.commit()
            return r
        finally:
            c.close()

    def stop(self):
        for bp in (self.dms,):
            try:
                bp.worker.stop(timeout=10)
            except Exception:
                pass


class PracticeClient:
    def __init__(self, h, uid, ip="10.1.1.1"):
        self.h, self.uid, self.ip = h, uid, ip
        self.c = h.app.test_client()
        self.headers = {"Authorization": f"Bearer {h._tok(uid)}"}
        self.env = {"REMOTE_ADDR": ip}
        self.member_id = None

    def _req(self, method, url, **kw):
        return getattr(self.c, method)("/api/practice" + url if url.startswith("/") and not url.startswith("/api/") else url,
                                       headers=self.headers, environ_overrides=self.env, **kw)

    def get(self, url, **kw):
        return self._req("get", url, **kw)

    def post(self, url, json=None, **kw):
        if "data" in kw:
            return self._req("post", url, **kw)
        return self._req("post", url, json=json if json is not None else {}, **kw)

    def patch(self, url, json=None):
        return self._req("patch", url, json=json or {})

    def put(self, url, json=None):
        return self._req("put", url, json=json or {})

    def delete(self, url):
        return self._req("delete", url)

    def upload(self, name, data, **fields):
        import io as _io
        fields = {k: str(v) for k, v in fields.items() if v is not None}
        return self._req("post", "/api/dms/upload", data={"file": (_io.BytesIO(data), name), **fields}, content_type="multipart/form-data")

    def ok(self, resp, status=None):
        assert resp.status_code == (status or 200), (resp.status_code, resp.get_json())
        return resp.get_json()


@pytest.fixture
def practice(tmp_path, monkeypatch):
    h = PracticeHarness(tmp_path, monkeypatch)
    yield h
    h.stop()


@pytest.fixture
def firm(practice):
    """A practice with one Senior Advocate (Asha), one Junior (Ravi), one Office Staff member (Meena) and a second Junior (Kiran)."""
    h = practice
    asha = h.person("Asha Rao")
    asha.ok(asha.post("/firm", {"name": "Rao & Associates"}), 201)
    team = {"asha": asha}
    for key, name, role in (("ravi", "Ravi Nair", "junior"), ("meena", "Meena Das", "staff"), ("kiran", "Kiran Shah", "junior")):
        r = asha.ok(asha.post("/members", {"name": name, "email": f"{key}@example.com", "role": role, "password": "TempPass123!"}), 201)
        p = PracticeClient(h, r["member"]["user_id"])
        p.member_id = r["member"]["id"]
        team[key] = p
    asha.member_id = asha.ok(asha.get("/me"))["member"]["id"]
    h.team = team
    return h
