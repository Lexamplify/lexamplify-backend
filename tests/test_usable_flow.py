"""
tests/test_usable_flow.py - the guided drafting path end to end (backend side):
sign-in required, no invented facts, strict brief, case data, upload errors, Word export formatting.
"""
import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest
from unittest.mock import patch

from tests.helpers_app import make_app


class _Client:
    def __init__(self, c):
        self._c = c

    def _csrf(self):
        ck = self._c.get_cookie("csrf_access_token")
        return {"X-CSRF-TOKEN": ck.value} if ck else {}

    def post(self, *a, **kw):
        kw["headers"] = {**self._csrf(), **kw.get("headers", {})}
        return self._c.post(*a, **kw)

    def __getattr__(self, n):
        return getattr(self._c, n)


@pytest.fixture()
def anon(tmp_path, monkeypatch):
    a = make_app(tmp_path, monkeypatch, "usable.db")
    with a.test_client() as c:
        yield c


@pytest.fixture()
def client(anon):
    r = anon.post("/api/auth/register", json={"email": "adv@example.com", "password": "TestPass123!", "name": "Adv"})
    assert r.status_code == 201, r.get_json()
    return _Client(anon)


NOTICE_FACTS = (
    "MATTER DETAILS\n- Sender: Sharma Textiles, 14 Ring Road, Surat\n- Addressee: Bharat Dyes Pvt Ltd, Plot 14, GIDC Vapi\n"
    "- Amount claimed: Rs. 4,75,000\n- Court: Delhi High Court, CS(COMM) 214/2026"
)


def test_draft_requires_sign_in(anon):
    r = anon.post("/api/documents/draft", json={"instructions": "Draft a legal notice"})
    assert r.status_code == 401
    assert r.get_json()["code"] == "AUTH_REQUIRED"


def test_extract_text_requires_sign_in(anon):
    r = anon.post("/api/contract/extract-text", data={"file": (io.BytesIO(b"x"), "a.pdf")}, content_type="multipart/form-data")
    assert r.status_code == 401


def test_draft_prompt_carries_facts_rule_and_matter_details(client):
    with patch("routes.document_routes.ask_groq", return_value="LEGAL NOTICE ...") as m:
        r = client.post("/api/documents/draft", json={"prompt": NOTICE_FACTS, "context": "MATTER FILE\nCase no: CS(COMM) 214/2026"})
    assert r.status_code == 200
    sent = " ".join(str(x) for x in (m.call_args.args + tuple(m.call_args.kwargs.values())))
    assert "FACTS RULE" in sent
    assert "Never invent" in sent
    assert "Bharat Dyes Pvt Ltd" in sent
    assert "CS(COMM) 214/2026" in sent


def test_draft_oversized_input_rejected(client):
    r = client.post("/api/documents/draft", json={"instructions": "x" * 13000})
    assert r.status_code == 413


def test_strict_brief_leaves_missing_facts_bracketed(client):
    qs = [
        {"id": "sender", "label": "Who is sending the notice?", "type": "text"},
        {"id": "amount", "label": "Amount claimed", "type": "number"},
        {"id": "cure", "label": "Days to comply", "type": "choice", "options": ["15 days", "30 days"]},
    ]
    r = client.post("/api/ai/intake/brief", json={"request": "Draft a legal notice", "questions": qs,
                                                  "answers": {"sender": "Sharma Textiles"}})
    assert r.status_code == 200
    brief = r.get_json()["brief"]
    assert "Sharma Textiles" in brief
    assert "[Bracketed Field: Amount claimed]" in brief
    assert "[Bracketed Field: Days to comply]" in brief
    # nothing invented for the blanks
    assert "Rs." not in brief and "₹" not in brief


def test_strict_brief_requires_sign_in(anon):
    r = anon.post("/api/ai/intake/brief", json={"request": "x", "questions": [], "answers": {}})
    assert r.status_code in (401, 422)


REALISTIC_HTML = """<h1 style="text-align:center">LEGAL NOTICE</h1>
<p style="text-align:right">Date: 17 October 2026</p>
<p><strong>To,</strong><br>M/s Bharat Dyes Pvt Ltd,<br>Plot 14, GIDC Vapi, Gujarat 396195</p>
<h2>Subject: Recovery of Rs. 4,75,000/- under Section 138 of the Negotiable Instruments Act, 1881</h2>
<p style="text-align:justify">Under instructions from my client <strong>Sharma Textiles</strong>, I serve this <em>notice</em>:</p>
<ol><li><p>That cheque no. 004512 dated 02.09.2026 was dishonoured.</p></li>
<li><p>That you pay within <u>15 days</u>.</p><ul><li>Interest at 12% p.a.</li></ul></li></ol>
<table><tr><th>Item</th><th>Amount</th></tr><tr><td>Cheque</td><td>4,75,000</td></tr></table>
<p>Advocate</p>"""


def test_word_export_keeps_headings_lists_tables_and_alignment(client):
    from docx import Document
    r = client.post("/api/contract/export-form-docx", json={"html": REALISTIC_HTML, "title": "Legal Notice"})
    assert r.status_code == 200
    d = Document(io.BytesIO(r.data))
    texts = [p.text for p in d.paragraphs]
    assert "LEGAL NOTICE" in texts
    assert any(t.startswith("Subject: Recovery") for t in texts)
    styles = [p.style.name for p in d.paragraphs]
    assert "List Number" in styles and "List Bullet 2" in styles
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "4,75,000"
    title = next(p for p in d.paragraphs if p.text == "LEGAL NOTICE")
    assert title.alignment == 1 and title.runs[0].bold
    assert d.styles["Normal"].font.name == "Times New Roman"


def test_word_export_rejects_empty(client):
    assert client.post("/api/contract/export-form-docx", json={"html": "  "}).status_code == 400
