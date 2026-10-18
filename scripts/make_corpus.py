"""Build the golden test corpus: tests/fixtures/indian_legal/ (synthetic documents, no real people or firms).

    python scripts/make_corpus.py

Deterministic: the same command always writes the same facts (file bytes may differ in PDF/DOCX timestamps).
Every name, number and amount is invented. manifest.json records what each file is and what the system must do with it,
including the ground-truth facts of the cheque-bounce packs, so tests can check extraction against the truth.
"""
import io
import json
import os
import random
import zipfile
from datetime import date, timedelta

import fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "tests", "fixtures", "indian_legal")
rnd = random.Random(1947)

FIRST = ["Arun", "Priya", "Ramesh", "Lakshmi", "Venkatesh", "Meena", "Suresh", "Anita", "Karthik", "Divya", "Mohammed", "Fatima", "Harpreet", "Sunita", "Rajiv", "Kavita"]
LAST = ["Kumar", "Raman", "Gupta", "Iyer", "Reddy", "Sharma", "Nair", "Pillai", "Singh", "Rao", "Khan", "Bose", "Menon", "Patel", "Joshi", "Das"]
BUSINESS = ["Traders", "Agencies", "Industries", "Enterprises", "Infra", "Textiles", "Foods", "Logistics", "Pharma", "Builders"]
CITIES = [("Chennai", "600017", "Tamil Nadu"), ("Coimbatore", "641001", "Tamil Nadu"), ("Hosur", "635109", "Tamil Nadu"), ("Bengaluru", "560001", "Karnataka"),
          ("Hyderabad", "500001", "Telangana"), ("Mumbai", "400001", "Maharashtra"), ("Delhi", "110001", "Delhi"), ("Pune", "411001", "Maharashtra")]
BANKS = ["State Bank of India", "HDFC Bank", "ICICI Bank", "Axis Bank", "Canara Bank", "Indian Bank"]
REASONS = ["Funds insufficient", "Payment stopped by drawer", "Account closed", "Exceeds arrangement"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]

manifest = []


def person():
    return f"{rnd.choice(FIRST)} {rnd.choice(LAST)}"


def firm():
    return f"M/s {rnd.choice(LAST)} {rnd.choice(BUSINESS)} Pvt Ltd"


def addr():
    c = rnd.choice(CITIES)
    return f"{rnd.randint(2, 98)}, {rnd.choice(['Gandhi', 'Nehru', 'Cross', 'Lake View', 'Temple', 'Market'])} Road, {c[0]} - {c[1]}", c


def fmt_date(d):
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"


def inr(n):
    s = str(int(n))
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts + [tail]) if parts else tail


def add(name, kind, expect, **extra):
    manifest.append({"file": name, "kind": kind, "expect": expect, **extra})


def write(name, data):
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, name), "wb") as fh:
        fh.write(data if isinstance(data, bytes) else data.encode("utf-8"))


def docx_bytes(paragraphs, tables=None):
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    for t in tables or []:
        tbl = d.add_table(rows=len(t), cols=len(t[0]))
        tbl.style = "Table Grid"
        for i, row in enumerate(t):
            for j, cell in enumerate(row):
                tbl.cell(i, j).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pdf_text_bytes(text, pages_min=1):
    import textwrap
    doc = fitz.open()
    lines = []
    for raw in text.split("\n"):
        lines.extend(textwrap.wrap(raw, 105) or [""])
    per_page = 48
    pages = max(pages_min, (len(lines) + per_page - 1) // per_page)
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        chunk = "\n".join(lines[i * per_page:(i + 1) * per_page])
        page.insert_textbox(fitz.Rect(56, 56, 540, 790), chunk, fontsize=10, fontname="helv")
    return doc.tobytes()


def pdf_scanned_bytes(text):
    """An image-only PDF (no text layer): what a phone scan looks like."""
    font = None
    for cand in (r"C:\Windows\Fonts\DejaVuSans.ttf", r"C:\Windows\Fonts\arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if os.path.isfile(cand):
            font = ImageFont.truetype(cand, 22)
            break
    img = Image.new("L", (1240, 1754), 255)
    draw = ImageDraw.Draw(img)
    y = 90
    for line in text.split("\n")[:60]:
        draw.text((90, y), line[:95], fill=20, font=font)
        y += 28
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    return doc.tobytes()


# ── cheque-bounce packs (ground truth recorded) ─────────────────────────────────────────────────
def s138_pack(i):
    payee, drawer = firm(), firm()
    (_a1, c1), (_a2, c2) = addr(), addr()
    cheque_date = date(2026, rnd.randint(1, 5), rnd.randint(1, 25))
    presented = cheque_date + timedelta(days=rnd.randint(3, 40))
    memo = presented + timedelta(days=rnd.randint(1, 4))
    amount = rnd.choice([75000, 125000, 250000, 475000, 880000, 1250000, 2150000])
    cheque_no = f"{rnd.randint(100000, 999999)}"
    reason = rnd.choice(REASONS)
    bank = rnd.choice(BANKS)
    drawn = f"{bank}, {c2[0]} Main Branch"
    facts = {"cheque_no": cheque_no, "cheque_amount": f"Rs. {inr(amount)}/-", "cheque_date": fmt_date(cheque_date), "return_reason": reason,
             "return_memo_date": fmt_date(memo), "presentation_date": fmt_date(presented), "drawn_bank": drawn}
    memo_text = (f"{bank.upper()}, {c2[0].upper()} MAIN BRANCH\nCHEQUE RETURN MEMO\n\nMemo dated {memo.strftime('%d/%m/%Y')}\n\n"
                 f"The cheque no. {cheque_no} dated {cheque_date.strftime('%d/%m/%Y')} for Rs. {inr(amount)}/- drawn on {drawn} in favour of {payee} "
                 f"was presented on {presented.strftime('%d/%m/%Y')} and is returned unpaid.\nReason for return: {reason}.\n\nFor {bank}\nManager\n")
    return payee, drawer, facts, memo_text


# ── builders ──────────────────────────────────────────────────────────────────────────────────
def build():
    # 1. s.138 packs: memos as text PDF (x3), plain text (x1), scanned (x1); a notice in docx
    for i in range(1, 4):
        payee, drawer, facts, memo = s138_pack(i)
        write(f"s138_return_memo_{i}.pdf", pdf_text_bytes(memo))
        add(f"s138_return_memo_{i}.pdf", "pdf", "ok", doc_type="notice_138", facts=facts, parties={"payee": payee, "drawer": drawer})
        notice = [f"LEGAL NOTICE", f"To {drawer}", f"Sub: Cheque no. {facts['cheque_no']} dated {facts['cheque_date']} for {facts['cheque_amount']}",
                  f"My client {payee} presented the cheque on {facts['presentation_date']}; it was returned with the remark {facts['return_reason']} (memo dated {facts['return_memo_date']}).",
                  "You are called upon to pay within fifteen (15) days of receipt of this notice."]
        write(f"s138_notice_{i}.docx", docx_bytes(notice))
        add(f"s138_notice_{i}.docx", "docx", "ok", doc_type="notice_138", facts={k: facts[k] for k in ("cheque_no", "cheque_amount", "cheque_date", "return_reason")})
    payee, drawer, facts, memo = s138_pack(4)
    write("s138_return_memo_4.txt", memo)
    add("s138_return_memo_4.txt", "txt", "ok", doc_type="notice_138", facts=facts)
    payee, drawer, facts, memo = s138_pack(5)
    write("s138_return_memo_5_scanned.pdf", pdf_scanned_bytes(memo))
    add("s138_return_memo_5_scanned.pdf", "pdf", "BLANK_DOCUMENT", doc_type="notice_138", note="image-only: the system must say it is scanned, never guess")

    # 2. agreements and court papers (docx + pdf + txt mix)
    for i in range(1, 4):
        landlord, tenant = person(), person()
        txt = (f"RENT AGREEMENT\nThis Agreement is made on {fmt_date(date(2026, 8, i))} between {landlord} (Landlord) and {tenant} (Tenant).\n"
               f"1. The Landlord lets the flat at {addr()[0]} for 11 months from {fmt_date(date(2026, 9, 1))}.\n"
               f"2. Rent is Rs. {inr(rnd.choice([18000, 25000, 42000]))} per month, due by the 5th of each month.\n"
               f"3. Security deposit is Rs. {inr(rnd.choice([54000, 75000, 126000]))}, refundable on vacating.\n4. Either party may end the tenancy on one month's notice.\n"
               f"5. The courts at {rnd.choice(CITIES)[0]} have jurisdiction.\n")
        if i == 1:
            write("rent_agreement_1.docx", docx_bytes(txt.split("\n")))
            add("rent_agreement_1.docx", "docx", "ok", doc_type="rent_agreement")
        elif i == 2:
            write("rent_agreement_2.pdf", pdf_text_bytes(txt))
            add("rent_agreement_2.pdf", "pdf", "ok", doc_type="rent_agreement")
        else:
            write("rent_agreement_3.txt", txt)
            add("rent_agreement_3.txt", "txt", "ok", doc_type="rent_agreement")
    for i in (1, 2):
        seller, buyer = person(), person()
        txt = (f"SALE DEED\nThis Sale Deed is made at {rnd.choice(CITIES)[0]} on {fmt_date(date(2026, 9, 10 + i))} between {seller} (Vendor) and {buyer} (Purchaser).\n"
               f"The Vendor sells Plot No. {rnd.randint(1, 90)}, Survey No. {rnd.randint(10, 400)}/{rnd.randint(1, 9)}, measuring {rnd.choice([1200, 2400, 3600])} sq. ft., "
               f"for Rs. {inr(rnd.choice([3500000, 4500000, 9800000]))}.\nThe Vendor acquired it by sale deed registered as Document No. {rnd.randint(100, 999)}/{rnd.randint(2005, 2020)}.\n")
        write(f"sale_deed_{i}.docx" if i == 1 else f"sale_deed_{i}.pdf", docx_bytes(txt.split("\n")) if i == 1 else pdf_text_bytes(txt))
        add(f"sale_deed_{i}.docx" if i == 1 else f"sale_deed_{i}.pdf", "docx" if i == 1 else "pdf", "ok", doc_type="sale_deed")
    for i in (1, 2):
        pl, df = person(), firm()
        txt = (f"IN THE COURT OF THE DISTRICT JUDGE, {rnd.choice(CITIES)[0].upper()}\nO.S. No. ____ of 2026\n{pl} ... Plaintiff\nversus\n{df} ... Defendant\nPLAINT\n"
               f"1. The Plaintiff supplied goods worth Rs. {inr(rnd.choice([450000, 1250000]))} to the Defendant on {fmt_date(date(2025, 11, 4))}.\n"
               f"2. The Defendant has not paid despite repeated requests.\nThe Plaintiff prays for a decree for the said amount with costs.\n")
        write(f"plaint_{i}.docx" if i == 1 else f"plaint_{i}.pdf", docx_bytes(txt.split("\n")) if i == 1 else pdf_text_bytes(txt))
        add(f"plaint_{i}.docx" if i == 1 else f"plaint_{i}.pdf", "docx" if i == 1 else "pdf", "ok", doc_type="plaint_skeleton")
    for i in (1, 2):
        d = person()
        txt = f"AFFIDAVIT\nI, {d}, aged {rnd.randint(28, 70)} years, residing at {addr()[0]}, do hereby solemnly affirm and state on oath as follows:\n1. That I am a resident of the above address.\n2. That the contents are true.\n"
        write(f"affidavit_{i}.docx" if i == 1 else f"affidavit_{i}.txt", docx_bytes(txt.split("\n")) if i == 1 else txt)
        add(f"affidavit_{i}.docx" if i == 1 else f"affidavit_{i}.txt", "docx" if i == 1 else "txt", "ok", doc_type="affidavit")
    for i in (1, 2):
        a, b = firm(), firm()
        txt = (f"MUTUAL NON-DISCLOSURE AGREEMENT\nThis Agreement is made on {fmt_date(date(2026, 7, 1 + i))} between {a} and {b}.\n1. Purpose. Evaluating a joint bid.\n"
               f"2. Each party shall keep the other's confidential information secret.\n3. This Agreement binds the parties for {rnd.choice([2, 3, 5])} years.\n4. Governed by the laws of India.\n")
        write(f"nda_{i}.docx" if i == 1 else f"nda_{i}.pdf", docx_bytes(txt.split("\n")) if i == 1 else pdf_text_bytes(txt))
        add(f"nda_{i}.docx" if i == 1 else f"nda_{i}.pdf", "docx" if i == 1 else "pdf", "ok", doc_type="nda")
    for i in (1, 2):
        acc = person()
        txt = (f"IN THE COURT OF THE SESSIONS JUDGE, {rnd.choice(CITIES)[0].upper()}\nBAIL APPLICATION under Section 483 of the Bharatiya Nagarik Suraksha Sanhita, 2023\n{acc} ... Petitioner\n"
               f"versus\nThe State.\nCrime No. {rnd.randint(100, 900)}/2026 registered at {rnd.choice(CITIES)[0]} Police Station on {fmt_date(date(2026, 6, 2))}.\n"
               f"The Petitioner was arrested on {fmt_date(date(2026, 6, 5))} and is in judicial custody.\nGrounds: 1. False implication. 2. No previous record.\n")
        write(f"bail_application_{i}.docx" if i == 1 else f"bail_application_{i}.pdf", docx_bytes(txt.split("\n")) if i == 1 else pdf_text_bytes(txt))
        add(f"bail_application_{i}.docx" if i == 1 else f"bail_application_{i}.pdf", "docx" if i == 1 else "pdf", "ok", doc_type="bail_application")
    write("vakalatnama_1.docx", docx_bytes(["VAKALATNAMA", f"I, {person()}, appoint Adv. {person()} to appear for me in CC No. {rnd.randint(1, 500)}/2026."]))
    add("vakalatnama_1.docx", "docx", "ok", doc_type="vakalatnama")

    # 3. a 40-page commercial contract (known key dates and one clause risk to look for)
    clauses = [f"{n}. Clause {n}. The parties shall perform their obligations under Schedule {n % 7 + 1} diligently and keep proper records. "
               f"Payment under this clause is due within {rnd.choice([15, 30, 45])} days of invoice. " * 6 for n in range(1, 150)]
    clauses[2] = "3. Termination. The Client may terminate this Agreement at any time without notice and without paying for work done."
    clauses[7] = "8. Interest. Delay in payment attracts interest at 24% per month."
    big = "MASTER SERVICES AGREEMENT\nThis Agreement is made on 1 August 2026 between Sundaram Agencies Pvt Ltd and Kumar Infra LLP.\n\n" + "\n\n".join(clauses) + \
          "\n\n150. Term. This Agreement runs for 3 years from 1 August 2026 and renews unless notice is given 90 days before expiry.\n"
    write("commercial_contract_40pages.pdf", pdf_text_bytes(big, pages_min=40))
    write("commercial_contract_40pages.docx", docx_bytes(big.split("\n\n")))
    add("commercial_contract_40pages.pdf", "pdf", "ok", doc_type="contract", pages_at_least=40, key_dates=["1 August 2026", "90 days", "3 years"], risky=["without notice", "24% per month"])
    add("commercial_contract_40pages.docx", "docx", "ok", doc_type="contract", key_dates=["1 August 2026", "90 days", "3 years"], risky=["without notice", "24% per month"])
    write("scanned_contract.pdf", pdf_scanned_bytes("SERVICE AGREEMENT\nThe Client may terminate this Agreement at any time without notice.\n" * 20))
    add("scanned_contract.pdf", "pdf", "BLANK_DOCUMENT", doc_type="contract")

    # 4. Indic and mixed-script, messy tables
    write("mixed_script_notice.docx", docx_bytes(["சட்ட அறிவிப்பு / कानूनी नोटिस / LEGAL NOTICE", "அனுப்புநர்: குமார் டிரேடர்ஸ் (M/s Kumar Traders Pvt Ltd)", "प्रति: सुंदरम एजेंसीज, चेन्नई - 600017",
                                                   "Cheque no. 004817 dated 12/06/2026 for Rs. 4,75,000/- was returned unpaid. Reason: Funds Insufficient."]))
    add("mixed_script_notice.docx", "docx", "ok", doc_type="notice_138", facts={"cheque_no": "004817", "return_reason": "Funds Insufficient"})
    write("mixed_script_note.txt", "கடன் ஒப்பந்தம் / ऋण समझौता\nBorrower: अरुण कुमार (Arun Kumar)\nAmount: Rs. 2,50,000/-\n")
    add("mixed_script_note.txt", "txt", "ok")
    write("messy_tables_1.docx", docx_bytes(["SCHEDULE OF PAYMENTS", "see table"], tables=[[["Instalment", "Due date", "Amount"], ["1", "5 October 2026", "Rs. 1,25,000"], ["2", "5 November 2026", "Rs. 1,25,000"]],
                                                                                          [["Party", "Role"], ["Sundaram Agencies Pvt Ltd", "Seller"], ["Kumar Infra LLP", "Buyer"]]]))
    add("messy_tables_1.docx", "docx", "ok")
    write("messy_tables_2.docx", docx_bytes(["", "   ", "TERMS", "\t\t1.\tPayment on delivery.\n\n\n", "<<<scan artefacts>>> ~~~ ||| ", "Contact: 98765 43210 / sundaram@example.com"]))
    add("messy_tables_2.docx", "docx", "ok")

    # 5. hostile and broken files
    write("old_word_renamed.docx", bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 600)
    add("old_word_renamed.docx", "docx", "NOT_A_DOCX")
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Protected contract text for the password test.")
    write("password_protected.pdf", doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user"))
    add("password_protected.pdf", "pdf", "EXTRACTION_ERROR")
    write("empty.pdf", b"")
    add("empty.pdf", "pdf", "EMPTY_FILE")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "<w:document/>")
        z.writestr("word/vbaProject.bin", b"\x00macro")
    write("macro_enabled.docx", buf.getvalue())
    add("macro_enabled.docx", "docx", "MACRO_DOCUMENT")
    write("fake_pdf.pdf", "This is only text with a .pdf name.")
    add("fake_pdf.pdf", "pdf", "NOT_A_PDF")
    good = pdf_text_bytes("Truncated PDF source text that is long enough to have several objects.\n" * 30)
    write("truncated.pdf", good[: len(good) // 3])
    add("truncated.pdf", "pdf", "EXTRACTION_ERROR")

    with open(os.path.join(OUT, "manifest.json"), "w", encoding="utf-8") as fh:
        json.dump({"note": "Synthetic documents. No real people, firms or cases.", "files": manifest}, fh, indent=1, ensure_ascii=False)
    total = sum(os.path.getsize(os.path.join(OUT, f)) for f in os.listdir(OUT))
    print(f"{len(manifest)} documents, {total / 1024:.0f} KB in {OUT}")


if __name__ == "__main__":
    build()
