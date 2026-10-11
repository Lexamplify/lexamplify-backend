"""
utils/vault_docs_convert.py - turns documents into editable HTML and back, keeping the layout lawyers care about.

  docx_to_html(bytes)   Word file -> HTML (fonts, sizes, colours, alignment, indents, lists, tables, pictures, links)
  pdf_to_html(bytes)    text-based PDF -> HTML that keeps centred headings, right-aligned signatures, bold/italic/underline,
                        indents, justified paragraphs and ruled tables (it reads each word's position on the page)
  text_to_html(str)     plain text -> HTML
  html_to_docx(html)    editor HTML -> .docx (real Word lists, tables, pictures, fonts)
  html_to_pdf(html)     editor HTML -> .pdf (PyMuPDF's HTML engine; ReportLab when PyMuPDF is not installed)
  html_to_text(html)    editor HTML -> readable plain text (tables as rows, lists numbered)
No network, no LLM: pure conversion.
"""
import base64
import html as _html
import io
import re
from collections import Counter

DEFAULT_FONT = "Times New Roman"
DEFAULT_PT = 12.0
BLOCKS = ["p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "td", "th"]


def _esc(s):
    return _html.escape(s or "", quote=False)


def _attr(s):
    return _html.escape(s or "", quote=True)


# ── plain text <-> html ─────────────────────────────────────────────────────────────
def text_to_html(text):
    paras = [p.strip("\n") for p in re.split(r"\n\s*\n", (text or "").replace("\r\n", "\n")) if p.strip()]
    return "".join("<p>" + "<br>".join(_esc(line) for line in p.split("\n")) + "</p>" for p in paras)


def html_to_text(h):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(h or "", "html.parser")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for img in soup.find_all("img"):
        img.decompose()
    out = []

    def cell_text(c):
        return " ".join(x.strip() for x in c.get_text("\n").split("\n") if x.strip())

    def walk(parent, depth=0):
        for el in parent.children:
            name = getattr(el, "name", None)
            if not name:
                continue
            if name == "table":
                rows = []
                for tr in el.find_all("tr"):
                    cells = [cell_text(c) for c in tr.find_all(["td", "th"])]
                    if any(cells):
                        rows.append("  |  ".join(cells))
                if rows:
                    out.append("\n".join(rows))
            elif name in ("ul", "ol"):
                n = int(el.get("start") or 1) if name == "ol" else 0
                for li in el.find_all("li", recursive=False):
                    own = " ".join(p.get_text().strip() for p in li.find_all("p", recursive=False)) or "".join(
                        t for t in li.find_all(string=True, recursive=False)).strip()
                    mark = f"{n}. " if name == "ol" else "• "
                    n += 1
                    out.append("    " * depth + mark + own)
                    for sub in li.find_all(["ul", "ol"], recursive=False):
                        walk(li, depth + 1) if False else walk_list(sub, depth + 1)
            elif name in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "pre"):
                t = el.get_text().strip()
                if t:
                    out.append(t)
            elif name == "hr":
                out.append("-" * 40)
            elif name in ("div", "section", "article", "body"):
                walk(el, depth)
            else:
                t = el.get_text().strip()
                if t:
                    out.append(t)

    def walk_list(lst, depth):
        holder = BeautifulSoup("", "html.parser")
        walk(_Wrap(lst), depth)

    class _Wrap:
        def __init__(self, el):
            self.children = [el]

    walk(soup)
    if not out:
        return soup.get_text("\n").strip()
    return "\n\n".join(out).replace("\n\n    ", "\n    ").replace("\n\n• ", "\n• ")


# ── DOCX -> HTML ────────────────────────────────────────────────────────────────────
_ALIGN = {0: "left", 1: "center", 2: "right", 3: "justify"}


def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _para_attr(p, getter):
    v = getter(p.paragraph_format)
    if v is not None:
        return v
    for st in _style_chain(p.style):
        v = getter(st.paragraph_format)
        if v is not None:
            return v
    return None


def _run_font(r, p, getter):
    v = getter(r.font)
    if v is not None:
        return v
    if r.style is not None:
        for st in _style_chain(r.style):
            v = getter(st.font)
            if v is not None:
                return v
    for st in _style_chain(p.style):
        v = getter(st.font)
        if v is not None:
            return v
    return None


def _numbering_kinds(doc):
    """numId -> {ilvl: 'bullet'|'decimal'|...}"""
    kinds = {}
    try:
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        root = doc.part.numbering_part.element
        W = "{%s}" % ns["w"]
        abstract = {}
        for a in root.findall("w:abstractNum", ns):
            lv = {}
            for l in a.findall("w:lvl", ns):
                f = l.find("w:numFmt", ns)
                lv[int(l.get(W + "ilvl"))] = f.get(W + "val") if f is not None else "decimal"
            abstract[a.get(W + "abstractNumId")] = lv
        for n in root.findall("w:num", ns):
            ab = n.find("w:abstractNumId", ns)
            if ab is not None:
                kinds[int(n.get(W + "numId"))] = abstract.get(ab.get(W + "val"), {})
    except Exception:
        pass
    return kinds


def _css_len(emu):
    return f"{emu / 12700:.1f}pt"


def _color_hex(rgb):
    return "#%s" % str(rgb).lower()


def docx_to_html(data):
    from docx import Document
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    kinds = _numbering_kinds(doc)
    normal = doc.styles["Normal"] if "Normal" in [s.name for s in doc.styles] else None
    base_pt = (normal.font.size.pt if normal is not None and normal.font.size else DEFAULT_PT)
    base_font = (normal.font.name if normal is not None and normal.font.name else None)
    counters = Counter()
    imgs_left = [60]

    def pic_html(run_el):
        out = ""
        for blip in run_el.xpath(".//a:blip"):
            rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if not rid or imgs_left[0] <= 0:
                continue
            try:
                part = doc.part.related_parts[rid]
                blob = part.blob
                if len(blob) > 3_000_000 or not part.content_type.startswith("image/") or "emf" in part.content_type or "wmf" in part.content_type:
                    continue
                imgs_left[0] -= 1
                w = ""
                ext = run_el.xpath(".//wp:extent")
                if ext:
                    w = f' width="{int(int(ext[0].get("cx")) / 9525)}"'
                out += f'<img src="data:{part.content_type};base64,{base64.b64encode(blob).decode()}"{w}>'
            except Exception:
                continue
        return out

    def run_html(r, p):
        text = r.text
        if not text:
            return pic_html(r._r)
        t = _esc(text).replace("\t", "&emsp;").replace("\n", "<br>")
        if _run_font(r, p, lambda f: f.bold):
            t = f"<strong>{t}</strong>"
        if _run_font(r, p, lambda f: f.italic):
            t = f"<em>{t}</em>"
        if _run_font(r, p, lambda f: f.underline):
            t = f"<u>{t}</u>"
        if _run_font(r, p, lambda f: f.strike):
            t = f"<s>{t}</s>"
        if _run_font(r, p, lambda f: f.subscript):
            t = f"<sub>{t}</sub>"
        if _run_font(r, p, lambda f: f.superscript):
            t = f"<sup>{t}</sup>"
        css = []
        fam = _run_font(r, p, lambda f: f.name) or base_font
        if fam:
            css.append(f"font-family:{fam}")
        sz = _run_font(r, p, lambda f: f.size)
        css.append(f"font-size:{(sz.pt if sz else base_pt):g}pt")
        try:
            col = _run_font(r, p, lambda f: f.color.rgb if f.color is not None and f.color.type is not None else None)
            if col is not None and str(col).lower() not in ("000000", "auto"):
                css.append(f"color:{_color_hex(col)}")
        except Exception:
            pass
        try:
            hl = r.font.highlight_color
            if hl is not None:
                css.append("background-color:" + {"YELLOW": "#ffff00", "BRIGHT_GREEN": "#00ff00", "TURQUOISE": "#00ffff", "PINK": "#ff00ff",
                                                  "BLUE": "#0000ff", "RED": "#ff0000", "GRAY_25": "#c0c0c0"}.get(hl.name, "#ffff00"))
        except Exception:
            pass
        return f'<span style="{";".join(css)}">{t}</span>' if css else t

    def inner_html(p):
        out = ""
        try:
            items = list(p.iter_inner_content())
        except Exception:
            items = list(p.runs)
        for it in items:
            if hasattr(it, "url") and hasattr(it, "runs"):
                body = "".join(run_html(r, p) for r in it.runs)
                out += f'<a href="{_attr(it.url or "")}">{body}</a>' if it.url else body
            else:
                out += run_html(it, p)
        return out

    def block_style(p, in_list=False):
        css = []
        al = _para_attr(p, lambda f: f.alignment)
        if al is not None and int(al) in _ALIGN and int(al) != 0:
            css.append(f"text-align:{_ALIGN[int(al)]}")
        li = None if in_list else _para_attr(p, lambda f: f.left_indent)      # a list's own indent / hanging indent comes from the list itself
        if li:
            css.append(f"margin-left:{_css_len(li)}")
        fi = None if in_list else _para_attr(p, lambda f: f.first_line_indent)
        if fi:
            css.append(f"text-indent:{_css_len(fi)}")
        sb = _para_attr(p, lambda f: f.space_before)
        if sb:
            css.append(f"margin-top:{_css_len(sb)}")
        sa = _para_attr(p, lambda f: f.space_after)
        if sa:
            css.append(f"margin-bottom:{_css_len(sa)}")
        ls = _para_attr(p, lambda f: f.line_spacing)
        if isinstance(ls, float) and 0.8 <= ls <= 3:
            css.append(f"line-height:{ls:g}")
        return f' style="{";".join(css)}"' if css else ""

    def list_info(p):
        pPr = p._p.pPr
        numPr = pPr.numPr if pPr is not None else None
        if numPr is not None and numPr.numId is not None and int(numPr.numId.val) != 0:
            num = int(numPr.numId.val)
            lvl = int(numPr.ilvl.val) if numPr.ilvl is not None else 0
            fmt = kinds.get(num, {}).get(lvl, "decimal")
            return num, lvl, ("ul" if fmt == "bullet" else "ol")
        name = (p.style.name if p.style is not None else "") or ""
        if name.lower().startswith("list"):
            m = re.search(r"(\d)$", name)
            lvl = int(m.group(1)) - 1 if m else 0
            ind = _para_attr(p, lambda f: f.left_indent)
            if not m and ind:                                   # our own export marks a nested level by its indent (0.25in + 0.3in per level)
                lvl = max(0, min(int(round((ind / 914400 - 0.25) / 0.3)), 4))
            return (hash(name) % 10_000) + 100_000, lvl, ("ul" if "bullet" in name.lower() else "ol")
        return None

    def heading_tag(p):
        name = (p.style.name if p.style is not None else "") or ""
        m = re.match(r"heading (\d)", name, re.I)
        if m:
            return f"h{min(max(int(m.group(1)), 1), 4)}"
        return "h1" if name.lower() == "title" else None

    def paragraph_blocks(paragraphs):
        """list of HTML strings for a run of paragraphs, grouping list items into <ul>/<ol>."""
        out, stack, cur_num = [], [], None

        def close_to(depth):
            while len(stack) > depth:
                out.append("</li></%s>" % stack.pop())

        for p in paragraphs:
            inner = inner_html(p)
            li = list_info(p)
            if li is None:
                close_to(0)
                cur_num = None
                if not inner.strip() and "<img" not in inner:
                    out.append("<p></p>")
                    continue
                tag = heading_tag(p) or "p"
                out.append(f"<{tag}{block_style(p)}>{inner}</{tag}>")
                continue
            num, lvl, tag = li
            if cur_num is not None and num != cur_num and lvl == 0:
                close_to(0)
            if lvl == 0:
                cur_num = num
            depth = lvl + 1
            if len(stack) > depth:
                close_to(depth)
            if len(stack) == depth:
                if stack[-1] != tag:
                    close_to(depth - 1)
                else:
                    out.append("</li>")
            while len(stack) < depth:
                start = ""
                if tag == "ol" and len(stack) == 0 and counters[num] > 0:
                    start = f' start="{counters[num] + 1}"'
                out.append(f"<{tag}{start}>")
                stack.append(tag)
                if len(stack) < depth:
                    out.append("<li>")
            counters[num] += 1 if len(stack) == 1 else 0
            out.append(f"<li><p{block_style(p, True)}>{inner}</p>")
        close_to(0)
        return out

    def cell_html(cell):
        paras = [Paragraph(el, cell) for el in cell._tc.iterchildren() if el.tag.endswith("}p")]
        return "".join(paragraph_blocks(paras)) or "<p></p>"

    def table_html(t):
        grid_rows, spans = [], {}
        trs = list(t._tbl.tr_lst)
        out = ["<table><tbody>"]
        for ri, tr in enumerate(trs):
            out.append("<tr>")
            col = 0
            for tc in tr.tc_lst:
                span = tc.grid_span or 1
                vm = tc.vMerge if hasattr(tc, "vMerge") else None
                if vm == "continue":
                    col += span
                    continue
                rowspan = 1
                if vm == "restart":
                    for nxt in trs[ri + 1:]:
                        pos, hit = 0, None
                        for c2 in nxt.tc_lst:
                            if pos == col:
                                hit = c2
                                break
                            pos += c2.grid_span or 1
                        if hit is not None and getattr(hit, "vMerge", None) == "continue":
                            rowspan += 1
                        else:
                            break
                attrs = ""
                if span > 1:
                    attrs += f' colspan="{span}"'
                if rowspan > 1:
                    attrs += f' rowspan="{rowspan}"'
                try:
                    fill = tc.tcPr.xpath("./w:shd/@w:fill") if tc.tcPr is not None else []
                    if fill and fill[0].lower() not in ("auto", "ffffff"):
                        attrs += f' style="background-color:#{fill[0].lower()}"'
                except Exception:
                    pass
                out.append(f"<td{attrs}>{cell_html(_Cell(tc, t))}</td>")
                col += span
            out.append("</tr>")
        out.append("</tbody></table>")
        return "".join(out)

    parts, run = [], []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            run.append(Paragraph(child, doc))
            continue
        if run:
            parts.extend(paragraph_blocks(run))
            run = []
        if tag == "tbl":
            parts.append(table_html(Table(child, doc)))
            parts.append("<p></p>")
    if run:
        parts.extend(paragraph_blocks(run))
    html = "".join(parts)
    html = re.sub(r"^(<p></p>)+|(<p></p>)+$", "", html)
    return html


# ── PDF -> HTML (layout-aware) ──────────────────────────────────────────────────────
def _fmt_of(fontname):
    f = (fontname or "").lower()
    return (any(k in f for k in ("bold", "black", "heavy", "semibold", "-bd")),
            any(k in f for k in ("italic", "oblique", "-it", "slanted")))


def _family_of(fontname):
    f = (fontname or "").lower()
    if any(k in f for k in ("arial", "helvet", "calibri", "verdana", "tahoma", "sans", "segoe", "roboto", "open")):
        return "Arial"
    if any(k in f for k in ("courier", "mono", "consolas")):
        return "Courier New"
    return None


def pdf_to_html(data, max_pages=150):
    """Re-creates the page layout as editable HTML. Returns '' for scanned PDFs (no text layer) so the caller can use OCR text."""
    import pdfplumber
    items_all = []          # per page: list of dict items sorted top->bottom
    sizes = Counter()
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for pi, page in enumerate(pdf.pages[:max_pages]):
            tables, tboxes = [], []
            try:
                tables = [t for t in page.find_tables() if t.bbox]
                tboxes = [t.bbox for t in tables]
            except Exception:
                pass

            def in_table(w):
                cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
                return any(b[0] - 1 <= cx <= b[2] + 1 and b[1] - 1 <= cy <= b[3] + 1 for b in tboxes)

            try:
                words = page.extract_words(extra_attrs=["fontname", "size"], keep_blank_chars=False, x_tolerance=1.6, y_tolerance=2.5)
            except Exception:
                words = []
            words = [w for w in words if not in_table(w)]
            # underline: thin horizontal strokes just under a word
            strokes = []
            for l in list(page.lines) + list(page.rects):
                h = abs(l["bottom"] - l["top"])
                if h <= 1.6 and (l["x1"] - l["x0"]) > 4:
                    strokes.append((l["x0"], l["x1"], l["top"]))

            def underlined(w):
                wl = w["x1"] - w["x0"]
                for x0, x1, y in strokes:
                    if -0.45 * w["size"] <= y - w["bottom"] <= 3.5 and min(x1, w["x1"]) - max(x0, w["x0"]) >= 0.7 * wl:
                        return True
                return False

            words.sort(key=lambda w: (w["top"], w["x0"]))
            lines = []
            for w in words:
                placed = False
                if lines:
                    ln = lines[-1]
                    mid = (w["top"] + w["bottom"]) / 2
                    if ln["top"] - 1 <= mid <= ln["bottom"] + 1:
                        ln["words"].append(w)
                        ln["top"], ln["bottom"] = min(ln["top"], w["top"]), max(ln["bottom"], w["bottom"])
                        placed = True
                if not placed:
                    lines.append({"words": [w], "top": w["top"], "bottom": w["bottom"]})
            for ln in lines:
                ln["words"].sort(key=lambda w: w["x0"])
                ln["x0"], ln["x1"] = ln["words"][0]["x0"], max(w["x1"] for w in ln["words"])
                ln["size"] = max(w["size"] for w in ln["words"])
                ln["page"] = pi
                ln["kind"] = "line"
                ln["underlined"] = {id(w): underlined(w) for w in ln["words"]}
                for w in ln["words"]:
                    sizes[round(w["size"] * 2) / 2] += len(w["text"])
            for t in tables:
                try:
                    rows = t.extract()
                except Exception:
                    rows = []
                if rows and any(any(c for c in r if c) for r in rows):
                    lines.append({"kind": "table", "top": t.bbox[1], "bottom": t.bbox[3], "x0": t.bbox[0], "x1": t.bbox[2], "rows": rows, "page": pi, "size": 0})
            lines.sort(key=lambda x: x["top"])
            # page geometry: body column from the long lines
            real = [l for l in lines if l["kind"] == "line"]
            wide = [l for l in real if (l["x1"] - l["x0"]) > 0.55 * page.width] or real
            for l in lines:
                l["_pw"] = page.width
            if real:
                L = min(l["x0"] for l in wide)
                R = max(l["x1"] for l in wide)
            else:
                L, R = 0, page.width
            for l in lines:
                l["L"], l["R"] = L, R
            items_all.extend(lines)
    if not sizes or sum(sizes.values()) < 15:
        return ""
    body_pt = sizes.most_common(1)[0][0]
    out = []

    def line_html(ln, body_pt):
        spans, prev = [], None
        for w in ln["words"]:
            b, i = _fmt_of(w.get("fontname"))
            u = ln["underlined"].get(id(w), False)
            gap = (w["x0"] - prev["x1"]) if prev else 0
            sep = ""
            if prev is not None:
                if gap < w["size"] * 0.12:
                    sep = ""
                elif gap < max(w["size"] * 2.6, 14):
                    sep = " "
                else:
                    sep = "&emsp;" * min(int(gap // max(w["size"] * 2.2, 18)) + 1, 8)
            spans.append((sep, w["text"], b, i, u, w["size"], _family_of(w.get("fontname"))))
            prev = w
        html, k = "", 0
        while k < len(spans):
            j = k
            key = spans[k][2:7]
            seg = ""
            while j < len(spans) and spans[j][2:7] == key:
                seg += (spans[j][0] if j > k else "") + _esc(spans[j][1]) if j > k else _esc(spans[j][1])
                j += 1
            lead = spans[k][0]
            b, i, u, sz, fam = key
            t = seg
            if b:
                t = f"<strong>{t}</strong>"
            if i:
                t = f"<em>{t}</em>"
            if u:
                t = f"<u>{t}</u>"
            css = []
            if fam:
                css.append(f"font-family:{fam}")
            if abs(round(sz * 2) / 2 - DEFAULT_PT) >= 0.5:
                css.append(f"font-size:{round(sz * 2) / 2:g}pt")
            if css:
                t = f'<span style="{";".join(css)}">{t}</span>'
            html += lead + t
            k = j
        return html

    def all_bold(ln):
        return all(_fmt_of(w.get("fontname"))[0] for w in ln["words"])

    def align_of(ln):
        if ln.get("force_center"):
            return "center"
        W = max(ln["R"] - ln["L"], 1)
        cx, mid = (ln["x0"] + ln["x1"]) / 2, (ln["L"] + ln["R"]) / 2
        if ln["x0"] - ln["L"] > 0.06 * W and abs(cx - mid) <= 0.03 * W:
            return "center"
        if ln["R"] - ln["x1"] <= 0.025 * W and ln["x0"] - ln["L"] > 0.3 * W:
            return "right"
        return "left"

    for idx in range(len(items_all) - 1):
        a, b = items_all[idx], items_all[idx + 1]
        if a["kind"] == "line" and b["kind"] == "line" and a["page"] == b["page"] and align_of(a) == "left" and align_of(b) == "center":
            Wd = max(a["R"] - a["L"], 1)
            if abs((a["x0"] + a["x1"]) / 2 - (a["L"] + a["R"]) / 2) <= 0.04 * Wd and 0 < b["top"] - a["top"] < 2.2 * a["size"]:
                a["force_center"] = True
    paras, cur = [], None

    def flush():
        nonlocal cur
        if cur:
            paras.append(cur)
        cur = None

    prev = None
    # typical line pitch of body text, to tell a paragraph gap from an ordinary line break
    pitches = []
    last = None
    for it in items_all:
        if it["kind"] == "line" and last is not None and last["kind"] == "line" and last["page"] == it["page"]:
            d = it["top"] - last["top"]
            if 0 < d < 3 * max(it["size"], 1):
                pitches.append(d)
        last = it
    pitch = sorted(pitches)[len(pitches) // 2] if pitches else body_pt * 1.4
    for it in items_all:
        if it["kind"] == "table":
            flush()
            paras.append({"table": it["rows"]})
            prev = it
            continue
        al = align_of(it)
        W = max(it["R"] - it["L"], 1)
        brk = True
        if cur and prev is not None and prev["kind"] == "line":
            same_page = prev["page"] == it["page"]
            dy = it["top"] - prev["top"] if same_page else None
            gap_ok = (dy is not None and dy <= pitch * 1.32) or (not same_page and (prev["x1"] > prev["R"] - 0.1 * W))
            if cur["align"] == al and gap_ok:
                if al in ("center", "right"):
                    brk = abs(it["size"] - prev["size"]) > 1.5 or all_bold(it) != all_bold(prev) and al == "center" and False
                else:
                    indented = it["x0"] - it["L"] > 0.02 * W and abs(it["x0"] - cur["lines"][-1]["x0"]) > 0.02 * W
                    short_prev = prev["x1"] < prev["R"] - 0.16 * W
                    bold_flip = all_bold(it) != all_bold(prev)
                    brk = indented or short_prev or bold_flip
        if brk:
            flush()
            cur = {"align": al, "lines": [it]}
        else:
            cur["lines"].append(it)
        prev = it
    flush()

    for p in paras:
        if "table" in p:
            rows = []
            for r in p["table"]:
                cells = "".join("<td><p>%s</p></td>" % "<br>".join(_esc(x.strip()) for x in (c or "").split("\n") if x.strip()) for c in r)
                rows.append(f"<tr>{cells}</tr>")
            out.append("<table><tbody>%s</tbody></table><p></p>" % "".join(rows))
            continue
        lines = p["lines"]
        L, R = lines[0]["L"], lines[0]["R"]
        W = max(R - L, 1)
        css = []
        if p["align"] in ("center", "right"):
            css.append(f"text-align:{p['align']}")
            joiner = "<br>"
        else:
            joiner = " "
            body_lines = lines[:-1]
            if len(lines) >= 3 and body_lines and sum(1 for l in body_lines if l["x1"] >= R - 0.03 * W) >= 0.7 * len(body_lines):
                css.append("text-align:justify")
            min_x = min(l["x0"] for l in lines)
            if min_x - L > 0.02 * W:
                css.append(f"margin-left:{(min_x - L):.0f}pt")
            first_extra = lines[0]["x0"] - min_x
            if len(lines) > 1 and first_extra > 0.02 * W:
                css.append(f"text-indent:{first_extra:.0f}pt")
            elif len(lines) == 1 and min_x - L > 0.02 * W:
                pass
        gap_before = 0
        pieces = [line_html(l, body_pt) for l in lines]
        html = pieces[0]
        for a_, b_ in zip(pieces, pieces[1:]):
            html += ("" if (joiner == " " and re.search(r"[A-Za-z]-(</[a-z]+>)*$", a_) and re.match(r"(<[a-z]+[^>]*>)*[a-z]", b_)) else joiner) + b_
        style = f' style="{";".join(css)}"' if css else ""
        out.append(f"<p{style}>{html}</p>")
    return "".join(out)


# ── HTML -> DOCX ────────────────────────────────────────────────────────────────────
def _len_pt(v):
    """css length -> points (px, pt, in, cm, em at 12pt)"""
    m = re.match(r"\s*(-?[\d.]+)\s*(px|pt|in|cm|mm|em|rem)?", v or "")
    if not m:
        return None
    n = float(m.group(1))
    u = m.group(2) or "px"
    return n * {"px": 0.75, "pt": 1, "in": 72, "cm": 28.346, "mm": 2.8346, "em": 12, "rem": 12}[u]


def _style_dict(el):
    out = {}
    for part in (el.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip()
    return out


def _rgb(v):
    v = (v or "").strip().lower()
    m = re.match(r"#([0-9a-f]{6})$", v)
    if m:
        return m.group(1)
    m = re.match(r"#([0-9a-f]{3})$", v)
    if m:
        return "".join(c * 2 for c in m.group(1))
    m = re.match(r"rgba?\((\d+)\s*,\s*(\d+)\s*,\s*(\d+)", v)
    if m:
        return "%02x%02x%02x" % tuple(min(int(x), 255) for x in m.groups())
    return None


def html_to_docx(html, title=None):
    from bs4 import BeautifulSoup, NavigableString, Tag
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches, Emu

    d = Document()
    st = d.styles["Normal"]
    st.font.name = DEFAULT_FONT
    st.font.size = Pt(DEFAULT_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    soup = BeautifulSoup(html or "", "html.parser")
    HL = {"ffff00": WD_COLOR_INDEX.YELLOW, "00ff00": WD_COLOR_INDEX.BRIGHT_GREEN, "00ffff": WD_COLOR_INDEX.TURQUOISE, "ff00ff": WD_COLOR_INDEX.PINK}

    def new_num(ordered_start=1):
        """a fresh numbering instance (restarts at 1) cloned from the template's 'List Number' list"""
        try:
            numbering = d.part.numbering_part.element
            base = d.styles["List Number"].element.pPr.numPr.numId.val
            abstract = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(base)][0].find(qn("w:abstractNumId")).get(qn("w:val"))
            ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
            nid = max(ids) + 1
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(nid))
            ab = OxmlElement("w:abstractNumId")
            ab.set(qn("w:val"), str(abstract))
            num.append(ab)
            ov = OxmlElement("w:lvlOverride")
            ov.set(qn("w:ilvl"), "0")
            so = OxmlElement("w:startOverride")
            so.set(qn("w:val"), str(ordered_start))
            ov.append(so)
            num.append(ov)
            numbering.append(num)
            return nid
        except Exception:
            return None

    def apply_run_fmt(r, f):
        r.bold = True if f.get("b") else None
        r.italic = True if f.get("i") else None
        r.underline = True if f.get("u") else None
        if f.get("s"):
            r.font.strike = True
        if f.get("sub"):
            r.font.subscript = True
        if f.get("sup"):
            r.font.superscript = True
        if f.get("font"):
            name = f["font"].split(",")[0].strip().strip("'\"")
            r.font.name = name
            rpr = r._r.get_or_add_rPr()
            rpr.rFonts.set(qn("w:eastAsia"), name)
        if f.get("size"):
            r.font.size = Pt(round(f["size"] * 2) / 2)
        if f.get("color"):
            r.font.color.rgb = RGBColor.from_string(f["color"].upper())
        if f.get("hl"):
            r.font.highlight_color = HL.get(f["hl"], WD_COLOR_INDEX.YELLOW)

    def add_picture(para, src, width_px=None):
        m = re.match(r"data:image/[a-z+.-]+;base64,(.+)$", src or "", re.I | re.S)
        if not m:
            return
        try:
            blob = base64.b64decode(m.group(1))
            width = Inches(min((width_px or 400) / 96, 6.0))
            para.add_run().add_picture(io.BytesIO(blob), width=width)
        except Exception:
            pass

    def runs(node, para, fmt):
        for ch in node.children:
            if isinstance(ch, NavigableString):
                t = str(ch)
                if t:
                    t = t.replace("\u00a0", " ").replace("\u2003", "\t") if False else t
                    parts = re.split(r"(\t|\u2003)", t)
                    for part in parts:
                        if part in ("\t", "\u2003"):
                            para.add_run().add_tab()
                        elif part:
                            apply_run_fmt(para.add_run(part), fmt)
            elif isinstance(ch, Tag):
                if ch.name == "br":
                    para.add_run().add_break()
                    continue
                if ch.name == "img":
                    try:
                        w = int(float(ch.get("width"))) if ch.get("width") else None
                    except ValueError:
                        w = None
                    add_picture(para, ch.get("src"), w)
                    continue
                f = dict(fmt)
                if ch.name in ("strong", "b"):
                    f["b"] = True
                if ch.name in ("em", "i"):
                    f["i"] = True
                if ch.name == "u":
                    f["u"] = True
                if ch.name in ("s", "strike", "del"):
                    f["s"] = True
                if ch.name == "sub":
                    f["sub"] = True
                if ch.name == "sup":
                    f["sup"] = True
                if ch.name == "a":
                    f["u"] = True
                    f["color"] = f.get("color") or "0563c1"
                if ch.name == "mark":
                    f["hl"] = _rgb((_style_dict(ch).get("background-color") or ch.get("data-color") or "#ffff00")) or "ffff00"
                sd = _style_dict(ch)
                if "font-family" in sd:
                    f["font"] = sd["font-family"]
                if "font-size" in sd:
                    pt = _len_pt(sd["font-size"])
                    if pt:
                        f["size"] = pt
                if "color" in sd and _rgb(sd["color"]):
                    f["color"] = _rgb(sd["color"])
                if "background-color" in sd and _rgb(sd["background-color"]):
                    f["hl"] = _rgb(sd["background-color"])
                if "font-weight" in sd and sd["font-weight"] in ("bold", "700", "800", "900"):
                    f["b"] = True
                if "font-style" in sd and sd["font-style"] == "italic":
                    f["i"] = True
                if "text-decoration" in sd and "underline" in sd["text-decoration"]:
                    f["u"] = True
                runs(ch, para, f)

    def para_format(para, el):
        sd = _style_dict(el)
        al = sd.get("text-align")
        if al in ("center", "right", "justify"):
            para.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[al]
        pf = para.paragraph_format
        if "margin-left" in sd and _len_pt(sd["margin-left"]):
            pf.left_indent = Pt(_len_pt(sd["margin-left"]))
        if "text-indent" in sd and _len_pt(sd["text-indent"]):
            pf.first_line_indent = Pt(_len_pt(sd["text-indent"]))
        if "margin-top" in sd and _len_pt(sd["margin-top"]) is not None:
            pf.space_before = Pt(_len_pt(sd["margin-top"]))
        if "margin-bottom" in sd and _len_pt(sd["margin-bottom"]) is not None:
            pf.space_after = Pt(_len_pt(sd["margin-bottom"]))
        if "line-height" in sd:
            try:
                v = float(sd["line-height"])
                if 0.8 <= v <= 3:
                    pf.line_spacing = v
            except ValueError:
                pass

    def block(container, el, list_ctx=None):
        name = el.name
        if name in ("ul", "ol"):
            level = (list_ctx or {}).get("level", -1) + 1
            num = None
            if name == "ol":
                try:
                    start = int(el.get("start") or 1)
                except ValueError:
                    start = 1
                num = new_num(start)
            ctx = {"level": level, "kind": name, "num": num, "n": int(el.get("start") or 1) if name == "ol" else 0}
            for li in el.find_all("li", recursive=False):
                block(container, li, ctx)
            return
        if name == "li":
            ctx = list_ctx or {"level": 0, "kind": "ul", "num": None, "n": 0}
            first = True
            for ch in li_children(el):
                if isinstance(ch, Tag) and ch.name in ("ul", "ol"):
                    block(container, ch, ctx)
                    continue
                para = container.add_paragraph()
                src = ch if isinstance(ch, Tag) and ch.name in ("p", "h1", "h2", "h3", "h4", "h5", "h6") else None
                if first:
                    wired = False
                    if ctx["kind"] == "ul":
                        try:
                            para.style = d.styles["List Bullet"]
                            wired = True
                        except KeyError:
                            pass
                    elif ctx.get("num"):
                        try:
                            para.style = d.styles["List Number"]
                            pPr = para._p.get_or_add_pPr()
                            numPr = pPr.get_or_add_numPr()
                            numPr.get_or_add_numId().val = ctx["num"]
                            numPr.get_or_add_ilvl().val = 0
                            wired = True
                        except Exception:
                            wired = False
                    if not wired:
                        mark = "• " if ctx["kind"] == "ul" else f"{ctx['n']}. "
                        para.add_run(mark)
                    ctx["n"] += 1
                    para.paragraph_format.left_indent = Inches(0.25 + 0.3 * max(ctx["level"], 0) + (0.0 if wired else 0.25))
                    para.paragraph_format.first_line_indent = Inches(-0.25)
                else:
                    para.paragraph_format.left_indent = Inches(0.5 + 0.3 * max(ctx["level"], 0))
                runs(src if src is not None else wrap_inline(ch), para, {})
                if src is not None:
                    sd = _style_dict(src)
                    if sd.get("text-align") in ("center", "right", "justify"):
                        para_format(para, src)
                first = False
            return
        if name == "table":
            rows = el.find_all("tr")
            if not rows:
                return
            ncols = max(sum(int(td.get("colspan") or 1) for td in tr.find_all(["td", "th"], recursive=False)) for tr in rows)
            t = container.add_table(rows=len(rows), cols=max(ncols, 1))
            t.style = d.styles["Table Grid"]
            occupied = set()
            for ri, tr in enumerate(rows):
                ci = 0
                for td in tr.find_all(["td", "th"], recursive=False):
                    while (ri, ci) in occupied:
                        ci += 1
                    cs, rs = int(td.get("colspan") or 1), int(td.get("rowspan") or 1)
                    cell = t.cell(ri, ci)
                    if cs > 1 or rs > 1:
                        cell = cell.merge(t.cell(min(ri + rs - 1, len(rows) - 1), min(ci + cs - 1, ncols - 1)))
                        for rr in range(ri, ri + rs):
                            for cc in range(ci, ci + cs):
                                occupied.add((rr, cc))
                    fill = _rgb(_style_dict(td).get("background-color"))
                    if fill:
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:color"), "auto")
                        shd.set(qn("w:fill"), fill)
                        tcPr.append(shd)
                    inner = [c for c in td.children if isinstance(c, Tag)]
                    first = True
                    if not inner:
                        cell.paragraphs[0].add_run(td.get_text())
                    for c in inner:
                        if c.name in ("ul", "ol"):
                            block(cell, c)
                            continue
                        para = cell.paragraphs[0] if first and not cell.paragraphs[0].text and len(cell.paragraphs) == 1 else cell.add_paragraph()
                        first = False
                        para_format(para, c)
                        runs(c, para, {"b": True} if td.name == "th" else {})
                    ci += cs
            d.add_paragraph()
            return
        if name == "hr":
            para = d.add_paragraph()
            pPr = para._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bt = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
                bt.set(qn(k), v)
            bd.append(bt)
            pPr.append(bd)
            return
        if name == "img":
            para = container.add_paragraph()
            add_picture(para, el.get("src"), int(float(el.get("width"))) if el.get("width") else None)
            return
        if name in ("div", "section", "article", "body"):
            for ch in el.children:
                if isinstance(ch, Tag):
                    block(container, ch)
                elif isinstance(ch, NavigableString) and str(ch).strip():
                    container.add_paragraph(str(ch).strip())
            return
        if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            para = container.add_heading(level=min(int(name[1]), 4))
            for r in list(para.runs):
                r._r.getparent().remove(r._r)
            fmt = {"font": DEFAULT_FONT, "color": "000000", "b": True,
                   "size": {1: 18, 2: 15, 3: 13, 4: 12}[min(int(name[1]), 4)]}
        else:
            para = container.add_paragraph()
            fmt = {"b": True} if name == "th" else {}
            if name == "blockquote":
                para.paragraph_format.left_indent = Inches(0.5)
        para_format(para, el)
        runs(el, para, fmt)

    def li_children(li):
        kids = [c for c in li.children if not (isinstance(c, NavigableString) and not str(c).strip())]
        if kids and all(isinstance(c, NavigableString) or (isinstance(c, Tag) and c.name not in ("p", "ul", "ol")) for c in kids):
            return [li]            # bare inline text inside <li>
        return kids

    def wrap_inline(ch):
        w = BeautifulSoup("<span></span>", "html.parser").span
        if isinstance(ch, Tag) and ch is not None:
            return ch
        w.string = str(ch)
        return w

    top = [c for c in soup.children if isinstance(c, (Tag, NavigableString))]
    if not [c for c in top if isinstance(c, Tag)]:
        d.add_paragraph(soup.get_text())
    for el in top:
        if isinstance(el, Tag):
            block(d, el)
        elif str(el).strip():
            d.add_paragraph(str(el).strip())
    if title:
        d.core_properties.title = title[:200]
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ── HTML -> PDF ─────────────────────────────────────────────────────────────────────
_GENERIC = [
    (("arial", "helvet", "calibri", "verdana", "tahoma", "trebuchet", "segoe", "sans"), "sans-serif"),
    (("courier", "mono", "consolas"), "monospace"),
]


def _generic_family(name):
    n = (name or "").lower()
    for keys, g in _GENERIC:
        if any(k in n for k in keys):
            return g
    return "serif"


def _prep_pdf_html(html):
    """Normalises editor HTML for the PDF engines: generic font families, no scripts, images kept as data URIs."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html or "", "html.parser")
    for t in soup(["script", "style", "iframe", "object", "embed", "link"]):
        t.decompose()
    for el in soup.find_all(True):
        sd = _style_dict(el)
        if "font-family" in sd:
            sd["font-family"] = _generic_family(sd["font-family"])
        keep = {k: v for k, v in sd.items() if k in ("font-family", "font-size", "color", "background-color", "text-align", "margin-left", "text-indent",
                                                      "margin-top", "margin-bottom", "line-height", "font-weight", "font-style", "text-decoration")}
        if keep:
            el["style"] = ";".join(f"{k}:{v}" for k, v in keep.items())
        elif el.has_attr("style"):
            del el["style"]
        if el.name == "img" and not (el.get("src") or "").startswith("data:image/"):
            el.decompose()
    return str(soup)


PDF_CSS = """
body { font-family: serif; font-size: 12pt; line-height: 1.35; color: #000; }
p { margin: 0 0 7pt 0; }
h1 { font-size: 18pt; margin: 14pt 0 6pt 0; } h2 { font-size: 15pt; margin: 12pt 0 5pt 0; }
h3 { font-size: 13pt; margin: 10pt 0 4pt 0; } h4, h5, h6 { font-size: 12pt; margin: 8pt 0 4pt 0; }
ul, ol { margin: 0 0 7pt 0; } li { margin: 0 0 3pt 0; }
table { border-collapse: collapse; width: 100%; margin: 0 0 9pt 0; }
td, th { border: 0.75pt solid #444; padding: 3pt 5pt; vertical-align: top; } th { background-color: #eeeeee; font-weight: bold; }
blockquote { margin: 0 0 7pt 28pt; } hr { border: 0; border-top: 0.75pt solid #888; }
img { max-width: 100%; }
"""


def _pdf_via_pymupdf(html):
    try:
        import pymupdf as fz
    except ImportError:
        import fitz as fz
    story = fz.Story(html=f"<body>{_prep_pdf_html(html)}</body>", user_css=PDF_CSS)
    buf = io.BytesIO()
    writer = fz.DocumentWriter(buf)
    media = fz.paper_rect("a4")
    where = media + (64, 64, -64, -64)
    more, guard = 1, 0
    while more and guard < 400:
        dev = writer.begin_page(media)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
        guard += 1
    writer.close()
    return buf.getvalue()


def _pdf_via_reportlab(html, title=None):
    import os
    from bs4 import BeautifulSoup, NavigableString, Tag
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    unicode_ok = False
    serif, bold, ital, bital = "Times-Roman", "Times-Bold", "Times-Italic", "Times-BoldItalic"
    for base in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "C:/Windows/Fonts"):
        try:
            if os.path.exists(os.path.join(base, "DejaVuSerif.ttf")):
                for nm, fn in (("DVS", "DejaVuSerif.ttf"), ("DVS-B", "DejaVuSerif-Bold.ttf"), ("DVS-I", "DejaVuSerif-Italic.ttf"), ("DVS-BI", "DejaVuSerif-BoldItalic.ttf")):
                    pdfmetrics.registerFont(TTFont(nm, os.path.join(base, fn)))
                pdfmetrics.registerFontFamily("DVS", normal="DVS", bold="DVS-B", italic="DVS-I", boldItalic="DVS-BI")
                serif, unicode_ok = "DVS", True
                break
        except Exception:
            continue

    def clean(t):
        t = _esc(t)
        return t if unicode_ok else t.replace("₹", "Rs.").encode("latin-1", "replace").decode("latin-1")

    def inline(node, fmt=""):
        out = ""
        for ch in node.children:
            if isinstance(ch, NavigableString):
                out += clean(str(ch))
            elif isinstance(ch, Tag):
                if ch.name == "br":
                    out += "<br/>"
                    continue
                o, c = "", ""
                if ch.name in ("strong", "b"):
                    o, c = "<b>", "</b>"
                elif ch.name in ("em", "i"):
                    o, c = "<i>", "</i>"
                elif ch.name == "u":
                    o, c = "<u>", "</u>"
                elif ch.name in ("s", "strike"):
                    o, c = "<strike>", "</strike>"
                elif ch.name == "sub":
                    o, c = "<sub>", "</sub>"
                elif ch.name == "sup":
                    o, c = "<super>", "</super>"
                sd = _style_dict(ch)
                col = _rgb(sd.get("color"))
                if col:
                    o += f'<font color="#{col}">'
                    c = "</font>" + c
                out += o + inline(ch) + c
        return out

    story = []
    page_w = A4[0] - 2 * 56
    soup = BeautifulSoup(_prep_pdf_html(html), "html.parser")

    def style_for(el, base=12.0, name="body"):
        sd = _style_dict(el)
        size = base
        if el.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            size = {"h1": 18, "h2": 15, "h3": 13}.get(el.name, 12)
        al = {"center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY}.get(sd.get("text-align"), TA_LEFT)
        return ParagraphStyle(name, fontName=serif if el.name not in ("h1", "h2", "h3", "h4") else (bold if not unicode_ok else "DVS-B"),
                              fontSize=size, leading=size * 1.35, alignment=al,
                              leftIndent=(_len_pt(sd.get("margin-left")) or 0), firstLineIndent=(_len_pt(sd.get("text-indent")) or 0),
                              spaceAfter=(_len_pt(sd.get("margin-bottom")) if "margin-bottom" in sd and _len_pt(sd.get("margin-bottom")) is not None else 6),
                              spaceBefore=(_len_pt(sd.get("margin-top")) or 0))

    def emit(el, indent=0, mark=None):
        if isinstance(el, NavigableString):
            if str(el).strip():
                story.append(Paragraph(clean(str(el).strip()), style_for(soup.new_tag("p"))))
            return
        n = el.name
        if n in ("ul", "ol"):
            i = int(el.get("start") or 1)
            for li in el.find_all("li", recursive=False):
                m = "•" if n == "ul" else f"{i}."
                i += 1
                subs = li.find_all(["ul", "ol"], recursive=False)
                ps = li.find_all("p", recursive=False)
                text = "<br/>".join(inline(p) for p in ps) if ps else inline(li)
                if not ps:
                    for sl in subs:
                        sl.extract()
                    text = inline(li)
                st = style_for(li)
                st.leftIndent = indent + 18
                st.firstLineIndent = -14
                story.append(Paragraph(f"{m}&nbsp;&nbsp;{text}", st))
                for sl in subs:
                    emit(sl, indent + 18)
            return
        if n == "table":
            data, spans = [], []
            rows = el.find_all("tr")
            cols = max((len(r.find_all(["td", "th"])) for r in rows), default=1)
            for r in rows:
                row = []
                for td in r.find_all(["td", "th"]):
                    st = style_for(td, 10.5)
                    st.spaceAfter = 0
                    txt = "<br/>".join(inline(p) for p in td.find_all("p")) or inline(td)
                    row.append(Paragraph(("<b>%s</b>" % txt) if td.name == "th" else txt, st))
                row += [""] * (cols - len(row))
                data.append(row)
            if data:
                t = Table(data, colWidths=[page_w / cols] * cols, repeatRows=0)
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#444444")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
                story.append(t)
                story.append(Spacer(1, 8))
            return
        if n == "hr":
            story.append(Spacer(1, 4))
            return
        if n == "img":
            return
        if n in ("div", "section", "body", "html"):
            for ch in el.children:
                emit(ch, indent)
            return
        txt = inline(el)
        story.append(Paragraph(txt if txt.strip() else "&nbsp;", style_for(el)))

    for el in soup.children:
        emit(el)
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=56, rightMargin=56, topMargin=56, bottomMargin=56, title=(title or "")[:200]).build(story or [Spacer(1, 1)])
    return buf.getvalue()


def html_to_pdf(html, title=None):
    try:
        return _pdf_via_pymupdf(html)
    except ImportError:
        return _pdf_via_reportlab(html, title)
