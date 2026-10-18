"""Export a structured document (blocks) to Word and PDF. Output comes from the JSON blocks, not from HTML.

Word: A4, Times New Roman 12 pt, 1.5 line spacing, numbered paragraphs with hanging indent, page numbers,
[Bracketed Fields] highlighted yellow so nothing unresolved goes out unnoticed.
PDF: same blocks via fpdf2. It needs a Unicode TrueType font for the rupee sign; see _find_font().
"""
import io
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .render import BRACKET_RE

FONT = "Times New Roman"
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
          "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
MAX_BLOCKS = 800
MAX_TEXT = 6000


def _safe_blocks(blocks):
    """Accept only what the renderer produces; cap sizes so an export call cannot be used to exhaust memory."""
    if not isinstance(blocks, list) or len(blocks) > MAX_BLOCKS:
        raise ValueError("document too large or malformed")
    out = []
    for b in blocks:
        if not isinstance(b, dict) or not isinstance(b.get("type"), str):
            raise ValueError("malformed block")
        c = {"type": b["type"]}
        if isinstance(b.get("text"), str):
            c["text"] = b["text"][:MAX_TEXT]
        if isinstance(b.get("lines"), list):
            c["lines"] = [str(x)[:MAX_TEXT] for x in b["lines"][:60]]
        if isinstance(b.get("rows"), list):
            c["rows"] = [[str(x)[:600] for x in row[:12]] for row in b["rows"][:200] if isinstance(row, list)]
            c["header"] = bool(b.get("header"))
        for k in ("align",):
            if b.get(k) in _ALIGN:
                c[k] = b[k]
        if b.get("bold"):
            c["bold"] = True
        if isinstance(b.get("number"), int):
            c["number"] = b["number"]
        if isinstance(b.get("letter"), str):
            c["letter"] = b["letter"][:6]
        out.append(c)
    return out


def _runs(paragraph, text, bold=False, size=12, underline=False):
    """Add text, highlighting every [Bracketed Field: ...]."""
    pos = 0
    for m in BRACKET_RE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], bold, size, underline)
        r = _run(paragraph, m.group(0), bold, size, underline)
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], bold, size, underline)


def _run(paragraph, text, bold, size, underline):
    r = paragraph.add_run(text)
    r.bold = bold
    r.underline = underline
    r.font.name = FONT
    r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return r


def _fmt(p, align="left", space_after=8, line=1.5, left=None, first=None, keep_next=False):
    pf = p.paragraph_format
    p.alignment = _ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if left is not None:
        pf.left_indent = Cm(left)
    if first is not None:
        pf.first_line_indent = Cm(first)
    pf.keep_with_next = keep_next
    pf.widow_control = True


def _page_number_footer(section):
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Page ")
    r.font.name, r.font.size = FONT, Pt(10)
    for instr in ("PAGE",):
        run = p.add_run()
        run.font.name, run.font.size = FONT, Pt(10)
        for kind, text in (("begin", None), (None, instr), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)


def to_docx(blocks, title="Document", letterhead_lines=None) -> bytes:
    blocks = _safe_blocks(blocks)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.54)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.core_properties.title = str(title)[:120]
    doc.core_properties.author = "LexAmplify"
    _page_number_footer(sec)

    for line in (letterhead_lines or [])[:8]:
        p = doc.add_paragraph()
        _runs(p, str(line)[:200], bold=(line is letterhead_lines[0]), size=13 if line is letterhead_lines[0] else 11)
        _fmt(p, "center", 0, 1.0)

    for b in blocks:
        t = b["type"]
        if t == "pagebreak":
            doc.add_page_break()
        elif t == "spacer":
            _fmt(doc.add_paragraph(), space_after=6, line=1.0)
        elif t == "title":
            p = doc.add_paragraph()
            _runs(p, b.get("text", ""), bold=True, size=14, underline=True)
            _fmt(p, "center", 14, 1.5, keep_next=True)
        elif t == "heading":
            p = doc.add_paragraph()
            _runs(p, b.get("text", ""), bold=True)
            _fmt(p, "left", 6, 1.5, keep_next=True)
        elif t in ("right", "center"):
            p = doc.add_paragraph()
            _runs(p, b.get("text", ""), bold=bool(b.get("bold")))
            _fmt(p, t, 4, 1.15)
        elif t == "address":
            for i, line in enumerate(b.get("lines", [])):
                p = doc.add_paragraph()
                _runs(p, line, bold=(i == 0))
                _fmt(p, "left", 0 if i < len(b["lines"]) - 1 else 10, 1.15, left=0.0)
        elif t == "signature":
            for i, line in enumerate(b.get("lines", [])):
                p = doc.add_paragraph()
                _runs(p, line, bold=(i == 0))
                _fmt(p, "left", 0, 1.15, left=8.5, keep_next=(i < len(b["lines"]) - 1))
        elif t in ("numbered", "lettered"):
            p = doc.add_paragraph()
            label = f"{b.get('number', '')}." if t == "numbered" else b.get("letter", "")
            _runs(p, f"{label}\t{b.get('text', '')}")
            _fmt(p, "justify", 8, 1.5, left=1.27, first=-1.27)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
        elif t == "table":
            rows = b.get("rows", [])
            if rows:
                tbl = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                tbl.style = "Table Grid"
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        c = tbl.cell(ri, ci)
                        c.text = ""
                        _runs(c.paragraphs[0], cell, bold=(b.get("header") and ri == 0), size=11)
                _fmt(doc.add_paragraph(), space_after=6, line=1.0)
        else:  # para
            p = doc.add_paragraph()
            _runs(p, b.get("text", ""), bold=bool(b.get("bold")))
            _fmt(p, b.get("align", "left"), 8, 1.5, keep_next=bool(b.get("bold")))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_text(data: bytes) -> str:
    """Text of a .docx as the reader would see it (used to prove the file matches the preview)."""
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


# ── PDF ────────────────────────────────────────────────────────────────────────────────────────
def _has_rupee(path) -> bool:
    try:
        from fontTools.ttLib import TTFont
        f = TTFont(path, lazy=True, fontNumber=0)
        return 0x20B9 in f.getBestCmap()
    except Exception:  # noqa: BLE001 - an unreadable font simply is not chosen for the rupee sign
        return False


def _find_font():
    """Return (regular, bold, has_rupee). Prefers a font that can print the rupee sign.
    Order: DRAFT_FONT_DIR, static/fonts (bundled), then common system fonts."""
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    dirs = [os.environ.get("DRAFT_FONT_DIR", ""), os.path.join(root, "static", "fonts"),
            r"C:\Windows\Fonts", "/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/truetype/liberation",
            "/usr/share/fonts/truetype/liberation2", "/Library/Fonts", "/System/Library/Fonts/Supplemental"]
    pairs = [("DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf"), ("Tinos-Regular.ttf", "Tinos-Bold.ttf"),
             ("times.ttf", "timesbd.ttf"), ("LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf"),
             ("Times New Roman.ttf", "Times New Roman Bold.ttf"), ("arial.ttf", "arialbd.ttf")]
    fallback = None
    for d in dirs:
        if not d:
            continue
        for reg, bold in pairs:
            r, b = os.path.join(d, reg), os.path.join(d, bold)
            if os.path.isfile(r) and os.path.isfile(b):
                if _has_rupee(r):
                    return r, b, True
                fallback = fallback or (r, b, False)
    return fallback or (None, None, False)


def to_pdf(blocks, title="Document", letterhead_lines=None):
    """Returns (bytes, notes). `notes` lists anything the reader should know (e.g. fallback font used)."""
    from fpdf import FPDF

    blocks = _safe_blocks(blocks)
    notes = []
    reg, bold, has_rupee = _find_font()
    fonts = {}

    class _PDF(FPDF):
        def footer(self):
            self.set_y(-15)
            self.set_font(fonts.get("fam", "Times"), "", 10)
            self.cell(0, 8, f"Page {self.page_no()}", align="C")

    pdf = _PDF(format="A4", unit="mm")
    pdf.set_margins(25.4, 25.4, 25.4)
    pdf.set_auto_page_break(True, 25.4)
    pdf.set_title(str(title)[:120])
    if reg:
        pdf.add_font("Body", "", reg)
        pdf.add_font("Body", "B", bold)
        fonts["fam"] = "Body"
        fam = "Body"
        if has_rupee:
            subst = lambda s: s  # noqa: E731
        else:
            subst = lambda s: s.replace("₹", "Rs. ")  # noqa: E731
            notes.append("The PDF font has no rupee sign, so amounts are written as 'Rs.'. Install a bundled font (static/fonts) to fix this.")
    else:
        fam = "Times"
        notes.append("PDF used a basic font because no Unicode font is installed; the rupee sign is written as 'Rs.'.")
        subst = lambda s: s.replace("₹", "Rs. ").encode("latin-1", "replace").decode("latin-1")  # noqa: E731
    pdf.add_page()
    W = pdf.w - pdf.l_margin - pdf.r_margin

    def para(text, align="L", bold_all=False, size=12, h=7.0, indent=0.0, space=3.0):
        text = subst(text)
        pdf.set_x(pdf.l_margin + indent)
        pos = 0
        segs = []
        for m in BRACKET_RE.finditer(text):
            if m.start() > pos:
                segs.append((text[pos:m.start()], False))
            segs.append((m.group(0), True))
            pos = m.end()
        if pos < len(text):
            segs.append((text[pos:], False))
        if align in ("C", "R"):
            # centred/right-aligned lines are short: flatten to one string (brackets keep colour only in L/J)
            pdf.set_font(fam, "B" if bold_all else "", size)
            pdf.multi_cell(W - indent, h, text, align=align, new_x="LMARGIN", new_y="NEXT")
        else:
            for chunk, is_field in segs:
                pdf.set_font(fam, "B" if (bold_all or is_field) else "", size)
                pdf.set_text_color(160, 40, 30) if is_field else pdf.set_text_color(0, 0, 0)
                pdf.write(h, chunk)
            pdf.set_text_color(0, 0, 0)
            pdf.ln(h)
        pdf.ln(space)

    for line in (letterhead_lines or [])[:8]:
        para(str(line)[:200], "C", bold_all=True, size=12, h=6, space=0)
    for b in blocks:
        t = b["type"]
        if t == "pagebreak":
            pdf.add_page()
        elif t == "spacer":
            pdf.ln(4)
        elif t == "title":
            para(b.get("text", ""), "C", True, 14, 8, space=4)
        elif t == "heading":
            para(b.get("text", ""), "L", True)
        elif t == "right":
            para(b.get("text", ""), "R", bool(b.get("bold")), h=6, space=1)
        elif t == "center":
            para(b.get("text", ""), "C", bool(b.get("bold")), h=6, space=1)
        elif t == "address":
            for i, line in enumerate(b.get("lines", [])):
                para(line, "L", i == 0, h=6, space=0)
            pdf.ln(4)
        elif t == "signature":
            for i, line in enumerate(b.get("lines", [])):
                para(line, "L", i == 0, h=6, indent=W * 0.5, space=0)
        elif t in ("numbered", "lettered"):
            label = f"{b.get('number', '')}.  " if t == "numbered" else f"{b.get('letter', '')}  "
            para(label + b.get("text", ""), "L", indent=0.0 if t == "numbered" else 8.0)
        elif t == "table":
            for row in b.get("rows", []):
                para(" | ".join(row), "L", bool(b.get("header") and row is b["rows"][0]), size=11, h=6, space=1)
            pdf.ln(3)
        else:
            para(b.get("text", ""), "L", bool(b.get("bold")))
    out = bytes(pdf.output())
    return out, notes
